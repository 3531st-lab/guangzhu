"""CAD and spreadsheet views of the existing, approved engineering results."""

import hashlib
import json
from datetime import datetime, timedelta, timezone

from .. import electrical, project as storage


def write_deliverables(project_id, version, directory):
    import ezdxf
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    def read(relative):
        return json.loads(storage._path(project_id, relative).read_text(encoding="utf-8"))

    project = read("project.json")
    design = read(f"designs/v{version}.json")
    if project['roof']['type'] in ('ground', 'carport') and design.get('installation_roof'):
        project['roof'] = design['installation_roof']
    economy = read(f"analysis/v{version}/economics.json")
    strings = read(f"analysis/v{version}/electrical.json")
    structure = read(f"analysis/v{version}/structure.json")
    module = electrical.equipment_catalog()["modules"][project["equipment"]["module_id"]]
    string_label, inverter_label = electrical.labels(strings)
    racking = 'ballast_ew' if design['params'].get('layout') == 'ew' and project['roof']['type'] == 'flat' else project['equipment']['racking']
    racking_label = {'ballast_tilt': '压载斜支架', 'clamp_flush': '彩钢瓦夹具（示意）',
                     'clamp_barn': '养殖场耐氨腐蚀夹具及导轨（示意选型）',
                     'ballast_ew': '东西向压载支架（示意）', 'plane_flush': '逐面顺坡支架',
                     'ground_pile': '地面桩基支架（示意）', 'carport_steel': '车棚钢结构（示意）'}[racking]
    if racking == 'clamp_flush' and structure.get('wind'):
        racking_label = '彩钢瓦导轨连接件（示意）'
    panels = {panel["id"]: panel for panel in design["modules"]}
    roof = project["roof"]
    drawing = ezdxf.new("R2010", setup=["linetypes"])
    drawing.units = 6
    space = drawing.modelspace()
    for name in ("ROOF", "OBSTACLE", "KEEPOUT", "PV", "STRING", "TEXT"):
        drawing.layers.new(name, dxfattribs={"linetype": "DASHED" if name == "KEEPOUT" else "CONTINUOUS"})

    def polygon(points, layer):
        space.add_lwpolyline([point[:2] for point in points], close=True, dxfattribs={"layer": layer})

    polygon(roof["outline"], "ROOF")
    from .pitched import cad, sheets as pitched_sheets
    cad(roof, drawing, space)
    from .installation import cad as installation_cad, sheets as installation_sheets
    installation_cad(project, design, drawing, space)
    from .livestock import cad as livestock_cad, sheets as livestock_sheets
    livestock_cad(project, design, drawing, space)
    for obstacle in roof["obstacles"]:
        if obstacle["kind"] == "cylinder":
            space.add_circle(obstacle["center"], obstacle["radius_m"], dxfattribs={"layer": "OBSTACLE"})
        else:
            polygon(obstacle["footprint"], "OBSTACLE")
    for keepout in roof["keepouts"] + design.get("added_keepouts", []):
        polygon(keepout["polygon"], "KEEPOUT")
    for panel in panels.values():
        polygon(panel["corners"], "PV")
    for string in strings["strings"]:
        point = panels[string["module_ids"][0]]["corners"][0][:2]
        space.add_text(string["id"], dxfattribs={"layer": "STRING", "height": .25, "insert": point})
    params = design["params"]
    title = f"{project['name']} | v{version} | {design['statistics']['dc_kw']} kW | {len(panels)} modules"
    date = datetime.now(timezone(timedelta(hours=8))).isoformat(timespec="seconds")
    notes = [title, f"tilt {params['tilt_deg']} deg | azimuth {params['azimuth_deg']} deg | {date} | x East, y North"]
    if roof['type'] == 'pitched':
        notes[1] = f"Flush on each P# face; angles and material on face layers | {date} | x East, y North"
    for index, note in enumerate(notes):
        point = (min(p[0] for p in roof["outline"]), min(p[1] for p in roof["outline"]) - 1 - index)
        space.add_text(note, dxfattribs={"layer": "TEXT", "height": .4, "insert": point})
    dxf = directory / f"layout-v{version}.dxf"
    drawing.saveas(dxf)
    book = Workbook()
    book.remove(book.active)

    def sheet(name, rows):
        page = book.create_sheet(name)
        for row in rows:
            page.append(row)
        for cells in page:
            for cell in cells:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        page.freeze_panes = "A2"
        for cell in page[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="9A522C")
        for column in page.columns:
            page.column_dimensions[column[0].column_letter].width = min(55, max(14, max(len(str(c.value or "")) for c in column) + 2))

    sheet("材料清单", [
        ["材料", "厂家", "型号", "单块功率 W", "数量", "单位", "合计 kW / 总质量 kg", "说明 / 规格书"],
        ["组件", module["manufacturer"], module["model"], module["pmax_w"], len(panels), "块", design["statistics"]["dc_kw"], module["datasheet_url"]],
        ["逆变器", None, inverter_label, None, strings["inverter_count"], "台"],
        [racking_label if racking != "ballast_tilt" else "压载支架", None, racking_label, None, len(panels), "套", None, "示意"],
        ["组件、夹具及导轨总质量" if racking == "clamp_flush" else "组件、支架及压块总质量", None, None, None, None, "kg",
         structure["mass_kg"], "示意；含组件" if racking == "clamp_flush" else "示意；总质量含组件，不是仅支架压块重量"],
        ["组串", None, None, None, strings["string_count"], "串"],
        ["每串块数", None, string_label, None, strings["modules_per_string"], "块/串"],
    ] + [["直流电缆" if r['kind'] == 'dc' else "交流电缆", None,
          ('' if r['kind'] == 'dc' else '4×') + f"{r['section_mm2']} mm²", None, r['length_m'], 'm', None,
          '示意；敷设/温度现场复核；交流为四芯路线米，直流含1.1富余'] for r in strings.get('cables', {}).get('summary', [])])
    currents = {(row["inverter"], row["mppt"]): row for row in strings["mppt_currents"]}
    rows = [["串号", "块数", "逆变器", "MPPT", "最冷开路电压 V", "最热工作电压 V", "所属 MPPT 合计工作电流 A", "所属 MPPT 合计短路电流 A"]]
    for string in strings["strings"]:
        current = currents[string["inverter"], string["mppt"]]
        rows.append([string["id"], len(string["module_ids"]), string["inverter"], string["mppt"],
                     string.get('voc_cold_v', strings["voc_cold_v"] * len(string['module_ids']) / strings['modules_per_string']),
                     string.get('vmp_hot_v', strings["vmp_hot_v"] * len(string['module_ids']) /
                                strings.get('modules_per_string_short', strings['modules_per_string'])),
                     current["working_a"], current["short_a"]])
    sheet("组串表", rows)
    cable_rows = {r['id']: r for r in strings.get('cables', {}).get('dc', [])}
    if strings.get('cables', {}).get('routing'):
        sheet('串联布线', [['串号', '逆变器', 'MPPT', '串内顺序', '组件编号', '路径状态', '本串导体米（含富余）']] +
              [[s['id'], s['inverter'], s['mppt'], i + 1, mid, cable_rows[s['id']]['routing_status'],
                cable_rows[s['id']]['length_m'] if i == 0 else None]
               for s in strings['strings'] for i, mid in enumerate(s['module_ids'])])
    if strings.get('phases'):
        phase_rows = [['阶段', '材料', '型号', '数量', '单位']]
        for phase in (1, 2):
            members = [m for m in panels.values() if m.get('phase') == phase]
            units = [u for u in strings['per_inverter'] if u.get('phase') == phase]
            ids = {s['id'] for s in strings['strings'] if s.get('phase') == phase}
            phase_rows += [[phase, '组件', module['model'], len(members), '块'], [phase, '支架', racking_label, len(members), '套']]
            phase_rows += [[phase, '逆变器', u['model'], 1, '台'] for u in units]
            member_ids = {m['id'] for m in members}
            wind = structure.get('wind', {})
            wind_members = [r for r in wind.get('modules', []) if r['id'] in member_ids]
            for zone, detail in wind.get('zones', {}).items():
                count = sum(r['ballast_blocks'] for r in wind_members if r['zone'] == zone)
                if count:
                    phase_rows.append([phase, detail['label'] + '压块', f"{wind['assumptions']['block_kg']}kg/块", count, '块'])
            anchors = sum(r['anchors'] for r in wind_members if r.get('fixing') == 'anchor')
            if anchors:
                phase_rows.append([phase, '化学锚栓（含局部防水）', '示意5kN/个', anchors, '个'])
            if roof['type'] != 'pitched' and any('clamps' in r for r in wind_members):
                phase_rows.append([phase, '抗风夹具', '全站实际布点按组件归期',
                                   sum(r['clamps']['adopted'] for r in wind_members if 'clamps' in r), '个'])
            if roof['type'] == 'pitched':
                phase_rows.pop(next(i for i in range(len(phase_rows) - 1, -1, -1)
                                    if phase_rows[i][0] == phase and phase_rows[i][1] == '支架'))
                for plane in roof['planes']:
                    selected = [r for r in wind_members if r['plane'] == plane['id']]
                    if not selected:
                        continue
                    label = '金属导轨' if plane['material'] == 'metal' else '混凝土固定导轨'
                    phase_rows.append([phase, f"P{plane['id']} {label}", '示意', len(selected), '套'])
                    if plane['material'] == 'concrete':
                        phase_rows.append([phase, f"P{plane['id']} 锚固件", '示意', sum(r['anchors'] for r in selected), '个'])
                    else:
                        phase_rows.append([phase, f"P{plane['id']} 抗风夹具", '实际布点', sum(r['clamps']['adopted'] for r in selected), '个'])
            for kind in ('dc', 'ac'):
                routes = [r for r in strings['cables'][kind] if r['id'] in (ids if kind == 'dc' else {'I' + str(u['inverter']) for u in units})]
                phase_rows += [[phase, kind + '电缆', str(r['section_mm2']) + 'mm²', r['length_m'], 'm'] for r in routes]
        sheet('分期材料', phase_rows)
        sheet('分期组串', [['串号', '阶段', '逆变器', '块数']] +
              [[r['id'], r['phase'], r['inverter'], len(r['module_ids'])] for r in strings['strings']])
        sheet('分期收益', [['阶段', '造价元', '回本年', '全期净收益元']] +
              [[r['name'], r['capex_cny'], r['payback_years'], r['profit_cny']] for r in economy.get('phases', [])])
        sheet('分期逐年现金流', [['年份', '一期净现金元', '两期净现金元', '二期增加净现金元']] +
              [[a['year'], a['net_cny'], b['net_cny'], c['net_cny']] for a, b, c in zip(economy['phase_cashflows']['first'],
                economy['cashflows'], economy['phase_cashflows']['increment'], strict=True)])
    from ..models import Design
    from ..orientation import groups
    orientations = groups(Design.model_validate(design))
    if len(orientations) > 1:
        sheet('朝向组', [['坡面', '倾角 °', '朝向 °', '组件数', '装机 kW', '组件编号']] +
              [[*identity, len(members), len(members) * module['pmax_w'] / 1000, ','.join(m.id for m in members)]
               for identity, members in orientations.items()])
    metrics = [("造价 元", "capex_cny"), ("回本 年", "payback_years"), ("动态回本 年", "dynamic_payback_years"),
               ("收益率 %", "irr_pct"), ("NPV 元", "npv_cny"), ("度电成本 元/kWh", "lcoe_cny_kwh")]
    flows = [("年", "year"), ("发电 kWh", "generation_kwh"), ("自用 kWh", "self_use_kwh"), ("上网 kWh", "export_kwh"),
             ("收入 元", "income_cny"), ("运维 元", "om_cny"), ("净现金流 元", "net_cny"), ("累计 元", "cumulative_cny"),
             ("折现 元", "discounted_cny"), ("累计折现 元", "cumulative_discounted_cny")]
    sheet("经济测算", [["指标", "原值"]] + [[label, economy[key]] for label, key in metrics] + [[]] +
          [[label for label, _ in flows]] + [[row[key] for _, key in flows] for row in economy["cashflows"]])
    headers = ["中文名", "数值", "单位", "状态", "来源标题", "链接", "访问日期"]
    sheet("假设与来源", [headers] + [[economy["params"][key] if field == "数值" else source.get(field, "未提供")
                                      for field in headers] for key, source in economy["source"].items()])
    from .financial import sheets as financial_sheets
    from .wind import materials as wind_materials
    wind_materials(structure, sheet, book)
    pitched_sheets(project, design, read(f'analysis/v{version}/shading.json'), economy, structure, sheet, book)
    financial_sheets(economy, sheet, book)
    installation_sheets(project, design, structure, economy, sheet, book)
    livestock_sheets(project, design, structure, economy, strings, sheet, book)
    xlsx = directory / f"bom-v{version}.xlsx"
    book.save(xlsx)
    return {path.name: hashlib.sha256(path.read_bytes()).hexdigest() for path in (dxf, xlsx)}
