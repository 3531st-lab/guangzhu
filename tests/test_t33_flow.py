"""Two real offline tool flows; fixed-clock replay retains every derived fingerprint."""
from datetime import datetime
from pathlib import Path
import hashlib
import json
import os
from shutil import copyfile
from unittest.mock import patch

import ezdxf
from openpyxl import load_workbook
import pytest
import requests

from guangzhu import brief, installation, mcp_actions, project as storage
from guangzhu.config import ROOT
from guangzhu.models import Design
from test_t32_flow import freeze_demo


@pytest.fixture(scope='module', params=['ground', 'carport'])
def installation_flow(request, tmp_path_factory):
    kind = request.param
    temporary = tmp_path_factory.mktemp('t33-' + kind)
    destination = os.environ.get('GUANGZHU_T33_CAPTURE')
    assets = (Path(destination) if destination else temporary / 'assets') / kind
    assets.mkdir(parents=True, exist_ok=True)
    clocks = {}
    recorded = os.environ.get('GUANGZHU_T33_CLOCKS')
    replay = json.loads(Path(recorded).read_text(encoding='utf-8')) if recorded else None
    original = storage._write
    def write(path, record, exclusive=False):
        if isinstance(record, Design):
            key = path.parent.parent.name + '/' + path.name
            if replay is not None:
                record.created_at = datetime.fromisoformat(replay[key].replace('Z', '+00:00'))
            clocks[key] = record.created_at.isoformat()
            (assets / 'clocks.json').write_text(json.dumps(clocks, indent=2), encoding='utf-8')
        return original(path, record, exclusive)
    identifier = 't33-' + kind
    with patch.object(storage, 'PROJECTS_DIR', temporary / 'projects'), patch.object(storage, '_write', write):
        with patch.object(requests.Session, 'request', side_effect=AssertionError('T33离线工具不得联网')):
            created = mcp_actions._new_project(identifier, kind + '-demo')
            assert created['roof_type'] == kind
            mcp_actions._set_requirements(identifier, '新建安装完整离线示意', {})
            generated = mcp_actions._generate_layout(identifier, {}, offline=True)
            initial = mcp_actions._analyze(identifier, 1, offline=True)
            compare = mcp_actions._compare_options(identifier, 1, offline=True)
            version = compare['recommended']['balanced']['version']
            summary = mcp_actions._analyze(identifier, version, offline=True)
            crosscheck = mcp_actions._crosscheck_yield(identifier, version, offline=True)
            pages = [crosscheck]
            while pages[-1].get('next_group_offset') is not None:
                pages.append(mcp_actions._crosscheck_yield(identifier, version, offline=True,
                                                          group_offset=pages[-1]['next_group_offset']))
            exported = mcp_actions._export(identifier, version, offline=True)
            assert exported['ok'] and summary['rules']['hard_failed'] == 0
            context = dict(project=storage.load(identifier).model_dump(mode='json'),
                           design=storage.load_design(identifier, version).model_dump(mode='json'),
                           version=version, summary=summary, compare=compare)
            for key in ('rules', 'economics', 'electrical', 'shading', 'structure'):
                context[key] = mcp_actions._read(identifier, f'analysis/v{version}/{key}.json')
            context.update(crosscheck=crosscheck, review=None, approval=None, sessions=[], verification={}, sources={},
                           versions=[], relaxations=[], history=[], run={}, requirements=[dict(text='新建安装完整离线示意')],
                           generated_at='2026-10-07T00:00:00+08:00', run_id='T33-offline-no-AI', project_id=identifier)
            from guangzhu.report.view import render
            contents = render(context)
            assert '未提供 AI 会话' in contents and '新建安装实体与工程量' in contents
            (assets / 'proposal.html').write_text(contents, encoding='utf-8')
            record, design = storage.load(identifier), storage.load_design(identifier, version)
            scene = dict(project=context['project'], design=context['design'], shading=context['shading'],
                         electrical=context['electrical'], wind=context['structure']['wind'],
                         facilities=installation.facilities(record, design))
            flow = dict(created=created, generated=generated, initial=initial, compare=compare,
                        final=summary, crosscheck=pages, exported=exported)
            for filename, value in [('context.json', context), ('scene.json', scene), ('tool-flow.json', flow)]:
                (assets / filename).write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')
            for filename in [*exported['files'], 'manifest.json']:
                copyfile(storage._path(identifier, f'exports/v{version}/{filename}'), assets / filename)
            freeze_demo(assets, kind + '-demo', identifier, version)
            yield dict(kind=kind, assets=assets, context=context, scene=scene, version=version)
            hashes = {p.relative_to(assets).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest()
                      for p in assets.rglob('*') if p.is_file() and p.name != 'capture-manifest.json'}
            (assets / 'capture-manifest.json').write_text(json.dumps(hashes, indent=2), encoding='utf-8')


def test_two_demos_normal_tools_actual_engineering_exports(installation_flow):
    data = installation_flow
    context, kind = data['context'], data['kind']
    workbook = load_workbook(data['assets'] / f"bom-v{data['version']}.xlsx", data_only=True)
    assert '基础实体' in workbook.sheetnames
    cad = ezdxf.readfile(data['assets'] / f"layout-v{data['version']}.dxf")
    assert 'PILE' in cad.layers and 'CANOPY' in cad.layers
    if kind == 'ground':
        assert '道路围栏土地' in workbook.sheetnames
        assert context['economics']['land_rent']['area_mu'] == pytest.approx(14.4)
    else:
        assert '车位及棚面' in workbook.sheetnames
        assert min(row['tilt_deg'] for row in context['compare']['options']) == 0
        for string in context['electrical']['strings']:
            assert len({mid.split('-')[0] for mid in string['module_ids']}) == 1
    assert '新建安装面' in brief.render_brief(dict(template=kind + '-demo', id='t33-brief', text='新建示意'))
    from t33_visual import capture
    capture(data)
