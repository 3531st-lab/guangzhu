"""Independent routes, shared construction, recurring land cash and snapshot contracts."""
import json
import math

from openpyxl import load_workbook
import pytest
import requests
from shapely import affinity
from shapely.geometry import Polygon

from guangzhu import analysis, costs, economics, electrical, installation, layout, mcp_actions, mcp_server, structure
from guangzhu import project as storage
from guangzhu.config import ROOT
from guangzhu.models import Roof, StructuredRequirements
from test_t33_installation import project


@pytest.mark.parametrize('kind', ['ground', 'carport'])
def test_phase_budget_cash_grid_actual_facilities_and_export_agree(tmp_path, monkeypatch, kind):
    monkeypatch.setattr(storage, 'PROJECTS_DIR', tmp_path / 'projects')
    monkeypatch.setattr(requests.Session, 'request', lambda *a, **k: pytest.fail('T33 remains offline'))
    identifier = 't33-phase-' + kind
    storage.init(ROOT / 'data/demo/projects/changzhou-demo', identifier)
    record = project(kind, width=36, depth=40, **(dict(direction_deg=0) if kind == 'carport' else {}))
    record.id = identifier
    budget = 90000 if kind == 'ground' else 150000
    record.requirements.structured = StructuredRequirements(phase1_budget_cny=budget)
    record.equipment.inverter_id = None
    storage.save(record)
    mcp_actions._generate_layout(identifier, {}, offline=True)
    summary = analysis.analyze(identifier, 1, offline=True)
    design = storage.load_design(identifier, 1)
    assert {m.phase for m in design.modules} == {1, 2}
    assert summary['rules']['hard_failed'] == 0
    full = mcp_actions._read(identifier, 'analysis/v1/structure.json')
    electric = mcp_actions._read(identifier, 'analysis/v1/electrical.json')
    cash = mcp_actions._read(identifier, 'analysis/v1/economics.json')
    first = design.model_copy(deep=True)
    first.modules = [m.model_copy(update=dict(phase=None)) for m in design.modules if m.phase == 1]
    first_load = structure.calculate(record, first)
    first_electric = electrical.calculate(record, first, offline=True)
    quote = costs.itemized(record, len(first.modules) * .55, first, first_electric, first_load)
    assert design.statistics['phasing']['achieved_cny'] == pytest.approx(quote['capex_cny'])
    assert cash['phases'][0]['capex_cny'] == pytest.approx(quote['capex_cny'])
    assert quote['capex_cny'] <= budget
    assert cash['phase_cashflows']['first'][0]['net_cny'] == pytest.approx(-quote['capex_cny'])
    assert cash['first_phase_facilities'] == json.loads(json.dumps(first_load['facilities']))
    grid = lambda item: next(p['point'] for p in item['cables']['positions'] if p['label'] == 'GRID')
    assert grid(electric) == grid(first_electric) == design.statistics['phasing']['grid_point']
    fixed = ('ground_road', 'ground_fence') if kind == 'ground' else ('carport_steel',)
    quantities = lambda result: {r['key']: r['quantity'] for r in result['extra_cost_items']}
    for key in fixed:
        assert quantities(full)[key] == quantities(first_load)[key]
    facts = full['facilities']
    if kind == 'ground':
        assert cash['phase_cashflows']['first'][1]['land_rent_cny'] == cash['land_rent']['annual_cny']
        assert all(r['land_rent_cny'] == 0 for r in cash['phase_cashflows']['increment'])
        saved = {r['id']: r for r in facts['piles']}
        assert all(p == saved[p['id']] for p in first_load['facilities']['piles'])
    else:
        assert first_load['facilities']['columns'] == facts['columns']
        assert first_load['facilities']['steel_kg'] == facts['steel_kg']
        assert cash['business_models'][2]['name'] == '场地合作（租金未计）'
        assert cash['business_models'][2]['owner_first_year_cny'] == 0
        assert all(row['rent'] == 0 for mode in cash['business_models'] for row in mode['accounts'])
        assert not any('9%' in note for note in cash['assumptions'])
    assert not any(row['key'] == 'refurbishment' for row in cash['costs']['items'])
    from guangzhu.report.deliverables import write_deliverables
    write_deliverables(identifier, 1, tmp_path)
    book = load_workbook(tmp_path / 'bom-v1.xlsx', data_only=True)
    rows = list(book['分期材料'].values)[1:]
    for key in fixed:
        assert sum(r[3] for r in rows if r[1] == key) == quantities(full)[key]
        assert all(r[0] == 1 for r in rows if r[1] == key)
    if kind == 'ground':
        assert sum(r[3] for r in rows if r[2] == '空地物理架') == facts['rack_count']
        assert sum(r[3] for r in rows if r[2] == 'piles') == len(facts['piles'])
    else:
        assert sum(r[3] for r in rows if r[2] == 'columns') == len(facts['columns'])
    book.close()
    from guangzhu.report.installation import sheets
    from openpyxl import Workbook
    empty_first = design.model_copy(deep=True)
    for panel in empty_first.modules:
        panel.phase = 2
    empty_book = Workbook()
    empty_book.active.title = '材料清单'
    empty_book.create_sheet('分期材料').append(['期', '名称', '类型', '数量', '单位'])
    def sheet(name, rows):
        page = empty_book.create_sheet(name)
        for row in rows:
            page.append(row)
    sheets(record.model_dump(), empty_first.model_dump(), full, cash, sheet, empty_book)
    assert all(row[0] == 2 for row in list(empty_book['分期材料'].values)[1:])
    empty_book.close()
    reply = mcp_server.call('gz_project_summary', dict(project_id=identifier))
    assert not reply.is_error
    assert len(reply.content[0].text.encode('utf-8')) + 128 <= 2048
    assert json.loads(reply.content[0].text)['roof_type'] == kind


def test_land_rent_independent_lcoe_zero_default_and_three_mode_cash_oracle():
    record = project(width=120, depth=80)
    record.economics.land_rent_cny_mu_year = 0
    record.economics.degradation_annual_pct = 0
    record.economics.discount_rate = .08
    energy = dict(p50=1000, p90=900)
    zero = economics.calculate(record, 1, energy, sensitivity=False)
    record.economics.land_rent_cny_mu_year = 100
    rent = economics.calculate(record, 1, energy, sensitivity=False)
    assert rent['land_rent']['annual_cny'] == pytest.approx(1440)
    assert rent['lcoe_cny_kwh'] - zero['lcoe_cny_kwh'] == pytest.approx(1.44)
    for before, after in zip(zero['business_models'], rent['business_models'], strict=True):
        assert after['pretax']['npv_cny'] - before['pretax']['npv_cny'] == pytest.approx(-15371.677711567545)
        assert all(row['rent'] == 0 and row['land_rent_cny'] == 1440 for row in after['accounts'])
    record.economics.land_rent_cny_mu_year = None
    default = economics.calculate(record, 1, energy, sensitivity=False)
    assert default['land_rent']['annual_cny'] == pytest.approx(14.4 * 800)


@pytest.mark.parametrize('rows,columns', [(1, 1), (1, 2), (2, 1), (2, 2)])
def test_rotated_parking_retains_clear_dimensions_and_row_direction(rows, columns):
    record = project('carport', width=24, depth=40, bay_rows=rows, column_rows=columns, direction_deg=0, facing='row')
    outline = affinity.rotate(Polygon(record.roof.outline), 32, origin=(0, 0))
    roof = record.roof.model_copy(deep=True)
    roof.outline = list(outline.exterior.coords)[:-1]
    roof.carport.direction_deg = 32
    roof.planes = []
    rebuilt = Roof.model_validate(dict(roof.model_dump(), **installation.parking(roof)))
    assert len(rebuilt.carport.bays) == len(record.roof.carport.bays)
    assert all(p.slope_azimuth_deg == 58 for p in rebuilt.planes)
    assert all(Polygon(b.polygon).area == pytest.approx(2.4 * 5.1) for b in rebuilt.carport.bays)
    assert len(rebuilt.carport.columns['1']) >= 2 * columns


def test_ground_ew_physical_racks_roads_and_height_enter_actual_scene():
    from guangzhu.shading_geometry import build_scene
    record = project(width=80, depth=80, road_every=2)
    design = layout.generate_layout(record, dict(layout='ew', lower_edge_m=1.2))
    facts = installation.facilities(record, design)
    assert len(facts['piles']) == 2 * facts['rack_count']
    assert all(len(p['modules']) <= 4 for p in facts['piles'])
    scene = build_scene(record, design)
    assert min(p[2] for m in scene.design.modules for p in m.corners) == pytest.approx(1.2)
    road_area = facts['road_area_m2']
    fewer = layout.generate_layout(record, dict(layout='ew', road_every=4))
    assert installation.facilities(record, fewer)['road_area_m2'] < road_area
    none = layout.generate_layout(record, dict(layout='ew', road_every=0))
    assert installation.facilities(record, none)['road_area_m2'] == 0


def test_ground_dc_routes_drop_to_zero_and_ac_allowance_is_explicit():
    record = project(width=36, depth=40)
    design = layout.generate_layout(record, dict(lower_edge_m=1.2))
    electric = electrical.calculate(record, design, offline=True)
    for circuit in electric['cables']['dc']:
        path = circuit['route_3d']
        assert circuit['route_m'] == pytest.approx(sum(math.dist(a, b) for a, b in zip(path, path[1:])), abs=.002)
        assert any(p[2] == 0 for p in path)
    for circuit in electric['cables']['ac']:
        assert circuit['length_m'] == pytest.approx(circuit['route_m'] + 5, abs=.002)


def test_actual_lower_edge_changes_obstacle_ray_and_pile_tops_meet_inclined_rails():
    import numpy as np
    from guangzhu.models import Obstacle
    from guangzhu.shading_geometry import build_scene, ray_hits, sun_vector
    hits = []
    for height in (.3, 1.2):
        record = project()
        design = layout.generate_layout(record, dict(lower_edge_m=height))
        sample = design.modules[0].corners[0]
        design.installation_roof.obstacles = [Obstacle(id='pole', label='试验杆', kind='cylinder',
                                                       center=(sample[0], sample[1] - 2), radius_m=.4, height_m=1)]
        actual = installation.facilities(record, design)
        for pile in actual['piles']:
            assert pile['rail'][0][2] == pytest.approx(pile['top_m'])
            assert pile['rail'][1][2] == pytest.approx(pile['top_m'])
        scene = build_scene(record, design)
        hits.append(ray_hits(scene, np.array([sample]), sun_vector(10, 180), np.array([0]))[0])
    assert hits == ['pole', '']
