"""Real offline tools and production WebGL UI; no AGH or paid model transport."""
import json
import os
from pathlib import Path
import shutil

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright
import pypdfium2 as pdfium
import requests

from guangzhu import api, layout, mcp_actions, project as storage
from guangzhu.config import API_ORIGIN, API_PORT, ROOT


def test_full_offline_presentation_and_recovery(tmp_path, monkeypatch):
    assets = Path(os.environ.get('GUANGZHU_T35_CAPTURE', str(tmp_path / 'captures')))
    assets.mkdir(parents=True, exist_ok=True)
    monkeypatch.setattr(requests.Session, 'request', lambda *a, **k: (_ for _ in ()).throw(AssertionError('Offline only')))
    monkeypatch.setattr(storage, 'PROJECTS_DIR', tmp_path / 'projects')
    monkeypatch.setattr(api, 'ROOT', tmp_path)
    monkeypatch.setattr(api, 'project_runs', lambda identifier: [])
    shutil.copytree(ROOT / 'data/demo/projects/grazing-demo', tmp_path / 'data/demo/projects/grazing-demo')
    pid, calls = 't35-demo', []
    def tool(name, value):
        calls.append(dict(tool=name, result=value))
        (assets / 'tool-flow.json').write_text(json.dumps(calls, ensure_ascii=False, indent=2), encoding='utf-8')
        return value
    tool('new_project', mcp_actions._new_project(pid, 'grazing-demo'))
    tool('set_requirements', mcp_actions._set_requirements(pid, '离线演示，使用真实组串和导出，不调用模型', {}))
    tool('generate_layout', mcp_actions._generate_layout(pid, {}, offline=True))
    summary = tool('analyze', mcp_actions._analyze(pid, 1, offline=True))
    cross = tool('crosscheck_yield', mcp_actions._crosscheck_yield(pid, 1, offline=True))
    assert summary['rules']['hard_failed'] == 0, summary['rules']
    exported = tool('export', mcp_actions._export(pid, 1, offline=True))
    assert exported['ok']
    record = storage.load(pid)
    design = storage.load_design(pid, 1)
    context = dict(project=record.model_dump(mode='json'), design=design.model_dump(mode='json'), summary=summary,
                   project_id=pid, run_id='T35-offline-no-AI', version=1, compare=None, crosscheck=cross,
                   facts={}, review=None, approval=None, sessions=[], verification={}, sources={}, versions=[],
                   relaxations=[], history=[], run={}, requirements=[], generated_at='2026-10-07T00:00:00+08:00')
    for name in ('rules', 'economics', 'electrical', 'shading', 'structure'):
        context[name] = mcp_actions._read(pid, f'analysis/v1/{name}.json')
    from guangzhu.report.view import render
    proposal = render(context)
    assert '串联布线方案' in proposal and '未提供 AI 会话' in proposal
    (assets / 'proposal.html').write_text(proposal, encoding='utf-8')
    (assets / 'context.json').write_text(json.dumps(context, ensure_ascii=False, indent=2), encoding='utf-8')
    for name in exported['files']:
        shutil.copyfile(storage._path(pid, f'exports/v1/{name}'), assets / name)
    book = load_workbook(assets / 'bom-v1.xlsx', data_only=True)
    rows = list(book['串联布线'].values)[1:]
    electric = context['electrical']
    assert sum(row[6] or 0 for row in rows) == sum(row['length_m'] for row in electric['cables']['dc'])
    assert [row[4] for row in rows] == [mid for s in electric['strings'] for mid in s['module_ids']]
    book.close()
    # A real unanalysed version exercises the supported missing-data state.
    pending = layout.generate_layout(record, dict(tilt_deg=10))
    pending.version = 2
    pending.parent_version = 1
    storage.save_design(pid, pending)
    result1, result2 = api.result_data(pid, 1), api.result_data(pid, 2)
    result1['report'] = dict(files=['proposal.pdf'], hint='离线生成，未提供AI会话或审批')
    backend = TestClient(api.app, base_url=f'http://127.0.0.1:{API_PORT}', headers={'origin':API_ORIGIN})
    script = (ROOT / 'desktop/web/app.js').read_text(encoding='utf-8') + '\nwindow.testApp={state,render,viewer:()=>scene};'
    bridge = ('export function setAghOrigin() {}\nexport class Bridge {constructor(){this.ready=true;this.models=[];}'
              'setAvailable(){} setTheme(){} send(){}}')
    canned = {'/api/runs':dict(runs=[]), '/api/status':dict(agh=dict(connected=False)),
              '/api/setup/status':dict(installation=dict(installed=False,id='development'), required=False,ready=True,steps=[])}
    opened, errors, console_errors, observed = [], [], [], []
    with sync_playwright() as automation:
        browser = automation.chromium.launch(channel='msedge', headless=True, args=['--enable-webgl', '--use-angle=swiftshader'])
        page = browser.new_page(viewport=dict(width=1440, height=900))
        page.on('pageerror', lambda error: errors.append(str(error)))
        page.on('console', lambda message: console_errors.append(message.text) if message.type == 'error' else None)
        def respond(route):
            path = route.request.url.split('offline.test', 1)[-1]
            bare = path.split('?', 1)[0]
            if bare == '/api/open':
                opened.append(route.request.post_data_json)
                route.fulfill(json=dict(ok=True))
            elif bare == f'/api/projects/{pid}/result':
                route.fulfill(json=result2 if 'version=2' in path else result1)
            elif bare in canned:
                route.fulfill(json=canned[bare])
            elif bare.startswith('/api/roofs'):
                reply = backend.get(path)
                route.fulfill(status=reply.status_code, body=reply.content, content_type='application/json')
            elif bare == '/app.js':
                route.fulfill(body=script, content_type='text/javascript')
            elif bare == '/bridge.js':
                route.fulfill(body=bridge, content_type='text/javascript')
            elif bare.startswith('/api/'):
                route.fulfill(json=dict(exists=False))
            else:
                file = ROOT / 'desktop/web' / ('index.html' if bare == '/' else bare.lstrip('/'))
                route.fulfill(path=str(file)) if file.is_file() else route.abort()
        page.route('**/*', respond)
        page.goto('http://offline.test/#new')
        page.wait_for_selector('[data-roof="grazing-demo"]')
        assert page.title() == '光筑'
        page.click('[data-roof="grazing-demo"]')
        page.evaluate("location.hash='#result/t35-demo?version=1'")
        page.wait_for_function('window.testApp?.viewer()?.stats().frames > 3')
        page.click('[data-scene-mode="wiring"]')
        selected = electric['strings'][0]['id']
        page.select_option('#wire-string', selected)
        page.wait_for_function('(id)=>window.testApp.viewer().stats().selectedString===id', arg=selected)
        assert f'MPPT {electric["strings"][0]["mppt"]}' in page.inner_text('#wire-details')
        assert page.is_hidden('#heat-legend')
        page.click('[data-action="scene-focus"]')
        page.wait_for_timeout(250)
        page.screenshot(path=str(assets / 'wiring-selected-light.png'))
        for mode in ('heat', 'string', 'phase', 'wind', 'light', 'wiring'):
            page.click(f'[data-scene-mode="{mode}"]')
        page.click('[data-action="scene-reset"]')
        page.select_option('#wire-string', '')
        page.uncheck('#wire-ac')
        page.check('#wire-ac')
        page.evaluate("window.testApp.state.theme='dark';document.querySelector('#settings-open').click()")
        page.click('[data-theme="dark"]')
        page.click('#settings-close')
        page.set_viewport_size(dict(width=1920, height=1080))
        page.screenshot(path=str(assets / 'wiring-all-dark.png'))
        for width, height in ((1440, 900), (1366, 768), (390, 844)):
            page.set_viewport_size(dict(width=width, height=height))
            assert page.evaluate('document.documentElement.scrollWidth <= innerWidth + 1')
            if width > 1000:
                assert page.evaluate('document.documentElement.scrollHeight <= innerHeight + 1')
        page.screenshot(path=str(assets / 'wiring-narrow.png'))
        page.set_viewport_size(dict(width=1440, height=900))
        # Hold the old response, complete a newer navigation, then release the old response.
        page.evaluate('''payload => {
          const original = window.fetch;
          window.fetch = (url, options) => String(url).endsWith('?version=2') ?
            new Promise(resolve => window.releaseOld = () => resolve(new Response(JSON.stringify(payload)))) : original(url, options);
          location.hash = '#result/t35-demo?version=2';
        }''', result2)
        page.wait_for_function('Boolean(window.releaseOld)')
        page.evaluate("location.hash='#result/t35-demo?version=1'")
        page.wait_for_function('window.testApp.viewer()?.stats().frames > 2 && window.testApp.state.result.version===1')
        page.evaluate('window.releaseOld()')
        page.wait_for_timeout(100)
        assert page.evaluate('window.testApp.state.result.version') == 1
        page.reload()
        page.wait_for_function('window.testApp?.viewer()?.stats().frames > 2')
        page.evaluate("location.hash='#result/t35-demo?version=2'")
        page.wait_for_function('window.testApp.state.result.version===2')
        page.click('[data-scene-mode="wiring"]')
        assert page.is_disabled('#wire-string') and '缺少线路数据' in page.inner_text('#wire-details')
        page.evaluate("location.hash='#result/t35-demo?version=1'")
        page.wait_for_function('window.testApp.state.result.version===1 && window.testApp.viewer()?.stats().frames > 2')
        page.evaluate("document.querySelector('#scene canvas').getContext('webgl2').getExtension('WEBGL_lose_context').loseContext()")
        page.wait_for_selector('#scene-retry:visible')
        page.click('#scene-retry')
        page.wait_for_function('window.testApp.viewer()?.stats().available && window.testApp.viewer().stats().frames > 2')
        page.click('[data-result-tab="文件"]')
        page.click('[data-action="open-report"]')
        assert opened[-1]['kind'] == 'report' and opened[-1]['project_id'] == pid
        observed.append(dict(url=page.url, title=page.title(), scene=page.evaluate('window.testApp.viewer().stats()')))
        assert not errors, errors
        assert not console_errors, console_errors
        report_page = browser.new_page()
        report_page.set_content(proposal)
        report_page.pdf(path=str(assets / 'proposal.pdf'), print_background=True)
        browser.close()
    document = pdfium.PdfDocument(str(assets / 'proposal.pdf'))
    texts = [document[i].get_textpage().get_text_range() for i in range(len(document))]
    assert not any('<span' in text or '&lt;' in text for text in texts)
    selected_page = next(i for i, text in enumerate(texts) if '串联布线方案' in text)
    document[selected_page].render(scale=1.5).to_pil().save(assets / 'proposal-wiring.png')
    document.close()
    (assets / 'ui-checks.json').write_text(json.dumps(dict(observed=observed, errors=errors, console_errors=console_errors,
        browser='Edge/Playwright; Browser plugin not available', viewports=[[1440,900],[1920,1080],[1366,768],[390,844]],
        pdf_pages=len(texts), pdf_wiring_page=selected_page+1, native_open='Request verified; actual OS launch left to Claude'),
        ensure_ascii=False, indent=2), encoding='utf-8')
