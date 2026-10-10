"""Each new demo follows real offline tools to exports and actual browser/PDF pixels."""
import hashlib
import json
import os
from pathlib import Path
from shutil import copyfile

from openpyxl import load_workbook
import pytest
import requests

from guangzhu import installation, mcp_actions, project as storage


@pytest.mark.parametrize('kind', ['barn', 'grazing'])
def test_real_offline_demo_export(tmp_path, monkeypatch, kind):
    monkeypatch.setattr(storage, 'PROJECTS_DIR', tmp_path / 'projects')
    monkeypatch.setattr(requests.Session, 'request', lambda *a, **k: pytest.fail('T34 flow must remain offline'))
    assets = Path(os.environ.get('GUANGZHU_T34_CAPTURE', str(tmp_path / 'assets'))) / kind
    assets.mkdir(parents=True, exist_ok=True)
    identifier = 't34-' + kind
    flow = []
    def save(name, value):
        flow.append(dict(tool=name, result=value))
        (assets / 'tool-flow.json').write_text(json.dumps(flow, ensure_ascii=False, indent=2), encoding='utf-8')
        return value
    save('new_project', mcp_actions._new_project(identifier, kind + '-demo'))
    save('set_requirements', mcp_actions._set_requirements(identifier, '畜牧光伏离线示意；实际用地与结构由当地核定', {}))
    save('generate_layout', mcp_actions._generate_layout(identifier, {}, offline=True))
    summary = save('analyze', mcp_actions._analyze(identifier, 1, offline=True))
    version = 1
    if kind == 'grazing':
        # Exercise the published spacing edit in the real tool pipeline even when G10 already passes.
        save('edit_design', mcp_actions._edit_design(identifier, 1, '加大排距保留更多地面受光',
             [dict(op='set_param', name='extra_row_gap_m', value=2)], offline=True))
        version = 2
        summary = save('analyze', mcp_actions._analyze(identifier, version, offline=True))
        old = storage.load_design(identifier, 1)
        new = storage.load_design(identifier, 2)
        assert new.statistics['row_spacing_m'] == pytest.approx(old.statistics['row_spacing_m'] + 2)
        assert [m.corners for m in new.modules] != [m.corners for m in old.modules]
    cross = save('crosscheck_yield', mcp_actions._crosscheck_yield(identifier, version, offline=True))
    assert summary['rules']['hard_failed'] == 0, summary['rules']
    exported = save('export', mcp_actions._export(identifier, version, offline=True))
    assert exported['ok']
    record, design = storage.load(identifier), storage.load_design(identifier, version)
    context = dict(project=record.model_dump(mode='json'), design=design.model_dump(mode='json'),
                   summary=summary, version=version, compare=None, crosscheck=cross)
    for key in ('rules', 'economics', 'electrical', 'shading', 'structure'):
        context[key] = mcp_actions._read(identifier, f'analysis/v{version}/{key}.json')
    context.update(review=None, approval=None, sessions=[], verification={}, sources={}, versions=[], relaxations=[], history=[], run={},
                   requirements=[dict(text='畜牧光伏完整离线示意')], generated_at='2026-10-07T00:00:00+08:00',
                   run_id='T34-offline-no-AI', project_id=identifier)
    from guangzhu.report.view import render
    contents = render(context)
    assert '未提供 AI 会话' in contents and context['structure']['livestock']['title'] in contents
    (assets / 'proposal.html').write_text(contents, encoding='utf-8')
    scene = dict(project=context['project'], design=context['design'], shading=context['shading'],
                 electrical=context['electrical'], wind=context['structure']['wind'], livestock=context['structure']['livestock'])
    if kind == 'grazing':
        scene['facilities'] = installation.facilities(record, design)
    else:
        from guangzhu.roofplane import scene_geometry
        scene['roof_geometry'] = scene_geometry(record.roof, [])
    for name, value in [('context',context), ('scene',scene)]:
        (assets / (name + '.json')).write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')
    for name in [*exported['files'], 'manifest.json']:
        copyfile(storage._path(identifier, f'exports/v{version}/{name}'), assets / name)
    book = load_workbook(assets / f'bom-v{version}.xlsx', data_only=True)
    assert '畜牧用途' in book.sheetnames
    book.close()
    from t34_visual import capture
    capture(dict(kind=kind, assets=assets, scene=scene))
    from test_t32_flow import freeze_demo
    freeze_demo(assets, kind + '-demo', identifier, version)
    hashes = {p.relative_to(assets).as_posix():hashlib.sha256(p.read_bytes()).hexdigest()
              for p in assets.rglob('*') if p.is_file() and p.name != 'capture-manifest.json'}
    (assets / 'capture-manifest.json').write_text(json.dumps(hashes, indent=2), encoding='utf-8')
