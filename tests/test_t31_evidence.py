"""Actual offline product geometry, WebGL scene, PDF rendering and exported quantities."""
import json
import os
import shutil

import pytest
import pypdfium2
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

from guangzhu import analysis, layout, mcp_actions, project as storage
from guangzhu.config import ROOT
from guangzhu.models import Requirements
from guangzhu.report.figures import plan
from guangzhu.report.view import render


def test_real_wind_scene_pdf_and_material_quantities(tmp_path, monkeypatch):
    monkeypatch.setattr(storage, 'PROJECTS_DIR', tmp_path / 'projects')
    assets = ROOT / 'reports/assets/T31/export' if os.environ.get('T31_CAPTURE') == '1' else tmp_path / 'assets'
    assets.mkdir(parents=True, exist_ok=True)
    record = storage.init(ROOT / 'data/demo/projects/changzhou-demo', 't31-evidence')
    record.name = 'T31离线验收厂房（示意输入，无AI）'
    record.requirements = Requirements()
    record.roof.outline = [(0, 0), (28, 0), (28, 18), (0, 18)]
    record.roof.height_m = 10
    record.roof.parapet_height_m = 0
    record.roof.obstacles = []
    record.roof.keepouts = []
    record.roof.inverter_positions = [(14, 9)]
    record.roof.load_capacity_kn_m2 = 3  # Explicit synthetic roof; original demo 0.5 rejection is tested separately.
    storage.save(record)
    design = layout.generate_layout(record, dict(tilt_deg=0, edge_setback_m=.1, walkway_every_n_rows=0))
    storage.save_design(record.id, design)
    summary = analysis.analyze(record.id, 1, offline=True)
    context = dict(project=record.model_dump(mode='json'), design=design.model_dump(mode='json'), summary=summary, version=1,
                   compare=None, facts={}, review=None, approval=None, sessions=[], requirements=[], verification={}, sources={},
                   versions=[], relaxations=[], history=[], run={}, generated_at='2026-10-06T00:00:00+08:00',
                   run_id='T31-offline-no-AI', project_id=record.id, crosscheck=None)
    context.update({name: mcp_actions._read(record.id, f'analysis/v1/{name}.json')
                    for name in ['electrical', 'shading', 'structure', 'rules', 'economics']})
    assert context['rules']['hard_failed'] == 0, [r for r in context['rules']['items'] if r['status'] == 'fail']
    result = mcp_actions._export(record.id, 1, offline=True)
    assert result['ok']
    exported = storage._path(record.id, 'exports/v1')
    for path in exported.iterdir():
        shutil.copyfile(path, assets / path.name)
    wind = context['structure']['wind']
    assert {r['zone'] for r in wind['modules']} == {'corner', 'edge', 'middle'}
    book = load_workbook(assets / 'bom-v1.xlsx', data_only=True)
    material = [row for row in book['材料清单'].iter_rows(min_row=2, values_only=True) if str(row[0]).endswith('压块')]
    assert sum(row[4] for row in material) == wind['ballast_blocks']
    assert sum(row[6] for row in material) == wind['ballast_mass_kg']
    ballast_cost = next(r for r in context['economics']['costs']['items'] if r['key'] == 'ballast')
    assert ballast_cost['quantity'] * 1000 == wind['ballast_mass_kg']
    assert sum(row[9] for row in material) == pytest.approx(ballast_cost['subtotal_cny'])
    assert sum(row[5] for row in list(book['风荷载'].values)[1:] if isinstance(row[5] if len(row) > 5 else None, int)) == wind['ballast_blocks']
    book.close()
    html = render(context)
    (assets / 'proposal.html').write_text(html, encoding='utf-8')
    scene = {key: context[key] for key in ('project', 'design', 'shading', 'electrical')}
    scene['wind'] = wind
    web = ROOT / 'desktop/web'
    source = '''<html lang="zh"><meta charset="utf-8"><link rel="stylesheet" href="/styles.css">
      <body style="padding:24px;background:#f4ece3"><h2>T31 风荷载 · 离线示意工程验收</h2><main id="content"></main>
      <script src="/vendor/three.min.js"></script><script src="/vendor/OrbitControls.min.js"></script>
      <script type="module">import {createScene} from '/scene.js';
      window.showScene=data=>{
      document.querySelector('#content').innerHTML='<div id="scene" style="height:730px"></div>';
      window.viewer=createScene(document.querySelector('#scene'));
      viewer.setData(data);
      viewer.setMode('wind');
      };
      </script></body></html>'''
    with sync_playwright() as automation:
        browser = automation.chromium.launch(channel='msedge', headless=True, args=['--enable-webgl', '--use-angle=swiftshader'])
        page = browser.new_page(viewport=dict(width=1440, height=960))
        errors = []
        page.on('pageerror', lambda error: errors.append(str(error)))
        def route(request):
            path = request.request.url.split('offline.test')[-1]
            if path == '/':
                request.fulfill(body=source, content_type='text/html')
            else:
                file = web / path.lstrip('/')
                request.fulfill(path=str(file)) if file.is_file() else request.abort()
        page.route('**/*', route)
        page.goto('http://offline.test/')
        page.wait_for_function('Boolean(window.showScene)')
        page.evaluate('showScene', scene)
        page.wait_for_function('window.viewer?.stats().frames > 3')
        assert page.evaluate('viewer.stats().panels') == len(design.modules)
        page.screenshot(path=str(assets / 'wind-3d.png'))
        for mode in ('phase', 'string', 'heat', 'light', 'wind'):
            page.evaluate('(mode) => viewer.setMode(mode)', mode)
        page.evaluate('viewer.dispose()')
        page.locator('#content').evaluate('(element, svg) => element.innerHTML=svg', str(plan(context, wind=True)))
        page.locator('#content svg').evaluate('svg => Object.assign(svg.style, {height:"780px", width:"100%"})')
        page.screenshot(path=str(assets / 'wind-plan.png'))
        page.set_content(html)
        page.pdf(path=str(assets / 'proposal.pdf'), print_background=True, prefer_css_page_size=True)
        assert not errors
        browser.close()
    document = pypdfium2.PdfDocument(assets / 'proposal.pdf')
    texts = [document[i].get_textpage().get_text_range() for i in range(len(document))]
    wind_page = next(i for i, text in enumerate(texts) if '风荷载与分区压块' in text)
    document[wind_page].render(scale=1.5).to_pil().save(assets / 'proposal-wind-page.png')
    source_page = next(i for i, text in enumerate(texts) if '风荷载依据与待核实项' in text)
    document[source_page].render(scale=1.5).to_pil().save(assets / 'proposal-wind-sources.png')
    assert '不能代替结构计算' in ''.join(texts)
    assert '缺GB50009' in ''.join(texts)
    document.close()
    (assets / 'context.json').write_text(json.dumps(context, ensure_ascii=False, indent=2), encoding='utf-8')
    (assets / 'scene.json').write_text(json.dumps(scene, ensure_ascii=False, indent=2), encoding='utf-8')
    (assets / 'verification.json').write_text(json.dumps(dict(pdf_pages=len(texts), wind_page=wind_page + 1,
        modules=len(design.modules), ballast_blocks=wind['ballast_blocks'], ballast_kg=wind['ballast_mass_kg'],
        ballast_cost_cny=ballast_cost['subtotal_cny'], actual_export=True, hard_failed=0), indent=2), encoding='utf-8')
