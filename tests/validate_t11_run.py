"""Read-only T11 acceptance: python tests/validate_t11_run.py RUN NEW_OUTPUT_DIR.

Use --report-only --report-dir PREVIEW_DIR for an explicitly partial PDF check.
This does not export, generate reports, start a service, or update any archive.
"""
import argparse
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
import hashlib
from html import unescape
import json
import math
from pathlib import Path
import re
import sys
import time

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from scripts.agh_export_run import SECRET, sanitize


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def snapshot(folder):
    registered = set()
    for metadata_path in folder.rglob("run.json"):
        registered.update(contained(metadata_path.parent, name) for name in read(metadata_path).get("artifacts", {}))
    for manifest in folder.rglob("manifest.json"):
        registered.update(contained(manifest.parent, name) for name in read(manifest).get("files", {}))
    hashes, ignored = {}, []
    for path in folder.rglob("*"):
        if not path.is_file():
            continue
        relative = path.relative_to(folder).as_posix()
        if re.fullmatch(r"~\$bom-v[1-9][0-9]*\.xlsx", path.name) and path.resolve() not in registered:
            ignored.append(dict(path=relative, reason="Unregistered Office temporary lock; not an approved or archived deliverable"))
        else:
            hashes[relative] = digest(path)
    return hashes, ignored


def contained(base, name):
    path = (base / name).resolve()
    if not path.is_relative_to(base.resolve()):
        raise ValueError("Requested path escapes its evidence directory")
    return path


def read(path):
    return json.loads(path.read_text(encoding="utf-8-sig"))


def same(actual, expected):
    if isinstance(expected, (int, float)) and not isinstance(expected, bool):
        return isinstance(actual, (int, float)) and math.isclose(actual, expected, rel_tol=1e-12, abs_tol=1e-9)
    return actual == expected


def points_match(actual, expected):
    actual, expected = list(actual), list(expected)
    return len(actual) == len(expected) and all(abs(p[axis] - q[axis]) <= .001 for p, q in zip(actual, expected) for axis in (0, 1))


def bounds(points):
    points = list(points)
    return [min(p[0] for p in points), min(p[1] for p in points), max(p[0] for p in points), max(p[1] for p in points)]


class Audit:
    def __init__(self, folder):
        self.folder = folder
        self.checks = []
        self.sources = {}

    def check(self, name, passed, **facts):
        self.checks.append(dict(name=name, status="pass" if passed else "fail", **facts))

    def skip(self, name, reason):
        self.checks.append(dict(name=name, status="skipped", reason=reason))

    def source(self, relative):
        path = contained(self.capture, relative)
        actual = digest(path)
        expected = self.capture_meta.get("artifacts", {}).get(relative)
        self.check("capture SHA: " + relative, actual == expected, actual=actual, expected=expected)
        self.sources[path.relative_to(self.folder).as_posix()] = dict(sha256=actual, bytes=path.stat().st_size)
        return read(path)

    def load(self):
        self.meta = read(self.folder / "run.json")
        verification = read(self.folder / "verification.json")
        exports = verification.get("verified_exports", [])
        self.version = exports[-1]["version"] if exports else self.meta.get("export_version")
        if type(self.version) is not int or self.version < 1:
            raise ValueError("No positive selected export version")
        self.capture = contained(self.folder, self.meta.get("capture", "."))
        self.capture_meta = read(self.capture / "run.json")
        self.project = self.source("project/project.json")
        self.design = self.source(f"project/designs/v{self.version}.json")
        self.electrical = self.source(f"project/analysis/v{self.version}/electrical.json")
        self.summary = self.source(f"project/analysis/v{self.version}/summary.json")
        self.economics = self.source(f"project/exports/v{self.version}/economics.json")
        self.manifest = self.source(f"project/exports/v{self.version}/manifest.json")
        self.export = contained(self.capture, f"project/exports/v{self.version}")
        checks = verification.get("checks", [])
        self.check("four verified stages", len(checks) == 4 and all(row.get("passed") is True for row in checks), stages=checks)
        self.check("selected version identity", self.design["version"] == self.manifest["version"] == self.version,
                   selected=self.version, design=self.design["version"], manifest=self.manifest["version"])

    def binary(self, suffix):
        name = f"{'layout' if suffix == 'dxf' else 'bom'}-v{self.version}.{suffix}"
        path = contained(self.export, name)
        actual = digest(path)
        original = self.manifest["files"].get(name)
        relative = path.relative_to(self.capture).as_posix()
        archived = self.capture_meta.get("artifacts", {}).get(relative)
        self.check("approved binary SHA: " + name, actual == original == archived,
                   actual=actual, manifest=original, capture_registry=archived)
        self.sources[path.relative_to(self.folder).as_posix()] = dict(sha256=actual, bytes=path.stat().st_size)
        return path

    def dxf(self):
        import ezdxf
        drawing = ezdxf.readfile(self.binary("dxf"))
        space = drawing.modelspace()
        roof = self.project["roof"]
        polygons = list(space.query('LWPOLYLINE[layer=="PV"]'))
        outlines = list(space.query('LWPOLYLINE[layer=="ROOF"]'))
        boxes = [row for row in roof["obstacles"] if row["kind"] != "cylinder"]
        circles = [row for row in roof["obstacles"] if row["kind"] == "cylinder"]
        keepouts = roof["keepouts"] + self.design.get("added_keepouts", [])
        expected = dict(ROOF=1, OBSTACLE=len(boxes) + len(circles), KEEPOUT=len(keepouts),
                        PV=len(self.design["modules"]), STRING=self.electrical["string_count"], TEXT=2)
        counts = {name: len(space.query(f'*[layer=="{name}"]')) for name in expected}
        self.check("DXF six layers and counts", counts == expected and set(expected) <= {row.dxf.name for row in drawing.layers},
                   actual=counts, expected=expected)
        self.check("DXF metres and dashed keepouts", drawing.units == drawing.header["$INSUNITS"] == 6
                   and drawing.layers.get("KEEPOUT").dxf.linetype == "DASHED")
        self.check("DXF roof outline and bounds within 1 mm", len(outlines) == 1 and outlines[0].closed
                   and points_match(outlines[0].get_points("xy"), roof["outline"]),
                   actual_bounds=bounds(outlines[0].get_points("xy")) if outlines else None, expected_bounds=bounds(roof["outline"]))
        self.check("DXF every PV corner within 1 mm", len(polygons) == len(self.design["modules"])
                   and all(row.closed and points_match(row.get_points("xy"), panel["corners"])
                           for row, panel in zip(polygons, self.design["modules"])))
        actual_boxes = list(space.query('LWPOLYLINE[layer=="OBSTACLE"]'))
        actual_circles = list(space.query('CIRCLE[layer=="OBSTACLE"]'))
        self.check("DXF obstacle types and geometry", len(actual_boxes) == len(boxes) and len(actual_circles) == len(circles)
                   and all(row.closed and points_match(row.get_points("xy"), box["footprint"]) for row, box in zip(actual_boxes, boxes))
                   and all(points_match([row.dxf.center], [circle["center"]]) and abs(row.dxf.radius - circle["radius_m"]) <= .001
                           for row, circle in zip(actual_circles, circles)))
        actual_keepouts = list(space.query('LWPOLYLINE[layer=="KEEPOUT"]'))
        self.check("DXF keepout polygons", len(actual_keepouts) == len(keepouts)
                   and all(row.closed and points_match(row.get_points("xy"), area["polygon"]) for row, area in zip(actual_keepouts, keepouts)))
        labels = list(space.query('TEXT[layer=="STRING"]'))
        panels = {row["id"]: row for row in self.design["modules"]}
        self.check("DXF string labels and position", [row.dxf.text for row in labels] == [row["id"] for row in self.electrical["strings"]]
                   and all(points_match([label.dxf.insert], [panels[string["module_ids"][0]]["corners"][0]])
                           for label, string in zip(labels, self.electrical["strings"])))

    def xlsx_cables(self, material_rows):
        rows = [row for row in material_rows if row[0] in ("直流电缆", "交流电缆")]
        if "cables" not in self.electrical:
            self.check("XLSX cable is uncalculated without invented numbers", len(rows) == 1
                       and rows[0][0] == "直流电缆" and rows[0][1:7] == (None,) * 6 and rows[0][7] == "本期未计算")
            return
        expected = []
        for row in self.electrical["cables"]["summary"]:
            name = {"dc": "直流电缆", "ac": "交流电缆"}[row["kind"]]
            section = ("" if row["kind"] == "dc" else "4×") + f"{row['section_mm2']} mm²"
            expected.append((name, None, section, None, row["length_m"], "m", None))
        actual = sorted((row[:7] for row in rows), key=repr)
        expected.sort(key=repr)
        self.check("XLSX cable rows equal archived quantities sections and units", len(actual) == len(expected)
                   and all(all(same(a, b) for a, b in zip(row, saved)) for row, saved in zip(actual, expected))
                   and all(type(row[4]) in (int, float) and "示意" in str(row[7]) for row in rows),
                   rows=len(actual), expected_rows=len(expected))

    def xlsx(self):
        import openpyxl
        book = openpyxl.load_workbook(self.binary("xlsx"), data_only=False)
        required = {"材料清单", "组串表", "经济测算", "假设与来源"}
        missing = sorted(required - set(book.sheetnames))
        self.check("XLSX four core sheets", not missing, actual=book.sheetnames, missing=missing)
        if missing:
            book.close()
            return
        material_rows = list(book["材料清单"].iter_rows(min_row=2, values_only=True))
        material = {row[0]: row for row in material_rows}
        metrics = [("组件", 4, len(self.design["modules"])), ("组件", 6, self.design["statistics"]["dc_kw"]),
                   ("压载支架", 4, len(self.design["modules"])), ("逆变器", 4, self.electrical["inverter_count"]),
                   ("组串", 4, self.electrical["string_count"]), ("每串块数", 4, self.electrical["modules_per_string"])]
        models = self.electrical.get("inverters")
        expected_models = [" + ".join(f"{row['model']} × {row['count']}" for row in models)] if models else [
            self.electrical["inverter_model"], f"{self.electrical['inverter_model']} × {self.electrical['inverter_count']}"]
        self.check("XLSX counts capacity and inverter model", all(same(material[name][column], value) for name, column, value in metrics)
                   and material["逆变器"][2] in expected_models
                   and same(material["组件"][4], self.electrical["module_count"]) and same(material["组件"][6], self.electrical["dc_kw"]),
                   electrical_module_count=self.electrical["module_count"], electrical_dc_kw=self.electrical["dc_kw"],
                   values=[dict(item=name, column=column + 1, actual=material[name][column], expected=value) for name, column, value in metrics])
        self.check("XLSX structural quantities marked indicative", "示意" in material["压载支架"][7]
                   and "示意" in material["组件、支架及压块总质量"][7] and "含组件" in material["组件、支架及压块总质量"][7])
        self.xlsx_cables(material_rows)
        structure = contained(self.capture, f"project/analysis/v{self.version}/structure.json")
        if structure.is_file():
            original = self.source(structure.relative_to(self.capture).as_posix())
            self.check("XLSX original structural mass", same(material["组件、支架及压块总质量"][6], original["mass_kg"]))
        else:
            self.skip("XLSX original structural mass", "Archive has no structure.json; indicative label checked, mass not independently compared")
        rows = list(book["组串表"].iter_rows(min_row=2, values_only=True))
        currents = {(row["inverter"], row["mppt"]): row for row in self.electrical["mppt_currents"]}
        expected = []
        longest = self.electrical["modules_per_string"]
        shortest = self.electrical.get("modules_per_string_short", longest)
        for string in self.electrical["strings"]:
            current = currents[string["inverter"], string["mppt"]]
            count = len(string["module_ids"])
            # Each string's own voltages: voc_cold_v is for the longest string and vmp_hot_v for the shortest one.
            expected.append([string["id"], count, string["inverter"], string["mppt"],
                             string.get("voc_cold_v", self.electrical["voc_cold_v"] * count / longest),
                             string.get("vmp_hot_v", self.electrical["vmp_hot_v"] * count / shortest),
                             current["working_a"], current["short_a"]])
        self.check("XLSX every string assignment voltage and MPPT current", len(rows) == len(expected)
                   and all(len(row) == len(original) and all(same(a, b) for a, b in zip(row, original))
                           for row, original in zip(rows, expected)), rows=len(rows))
        economy = list(book["经济测算"].iter_rows(values_only=True))
        keys = ["capex_cny", "payback_years", "dynamic_payback_years", "irr_pct", "npv_cny", "lcoe_cny_kwh"]
        self.check("XLSX six economic metrics", all(same(economy[index + 1][1], self.economics[key]) for index, key in enumerate(keys)))
        start = next(index for index, row in enumerate(economy) if row[0] == "年") + 1
        cash = economy[start:]
        keys = ["year", "generation_kwh", "self_use_kwh", "export_kwh", "income_cny", "om_cny", "net_cny",
                "cumulative_cny", "discounted_cny", "cumulative_discounted_cny"]
        self.check("XLSX 26 cashflow rows equal archived values", len(cash) == len(self.economics["cashflows"]) == 26
                   and [row[0] for row in cash] == list(range(26))
                   and all(all(same(row[index], original[key]) for index, key in enumerate(keys))
                           for row, original in zip(cash, self.economics["cashflows"])), rows=len(cash))
        sources = list(book["假设与来源"].iter_rows(min_row=2, values_only=True))
        expected = [[row["中文名"], self.economics["params"][key], row["单位"], row["状态"],
                     row["来源标题"] or None, row["链接"] or None, row["访问日期"] or None] for key, row in self.economics["source"].items()]
        self.check("XLSX all parameter values statuses and sources", len(sources) == len(expected)
                   and all(len(row) == len(original) and all(same(a, b) for a, b in zip(row, original))
                           for row, original in zip(sources, expected)),
                   indicative_expected=sum(row[3] == "示意" for row in expected), indicative_actual=sum(row[3] == "示意" for row in sources))
        self.check("XLSX no executable formulas", not any(cell.data_type == "f" for sheet in book for row in sheet for cell in row))
        book.close()

    def proposal(self, report_dir, preview=False):
        from pypdf import PdfReader
        html_path = contained(report_dir, f"report-v{self.version}.html")
        pdf_path = contained(report_dir, f"report-v{self.version}.pdf")
        document = html_path.read_text(encoding="utf-8")
        reader = PdfReader(pdf_path)
        pages = [page.extract_text() or "" for page in reader.pages]
        compact = lambda value: re.sub(r"\s+", "", value)
        sections = re.findall(r'<section[^>]*>(.*?)</section>', document, re.S)
        headings = [unescape(re.sub(r'<[^>]+>', '', re.search(r'<h1>(.*?)</h1>', section, re.S)[1])) for section in sections]
        positions = [next((index for index, page in enumerate(pages)
                           if any(compact(title) == compact(line) for line in page.splitlines())), -1) for title in headings]
        self.check("HTML 14 sections and PDF section headings", len(sections) == 14 and all(index >= 0 for index in positions)
                   and positions == sorted(set(positions)), html_sections=len(sections), pdf_pages=len(pages),
                   section_pages=[dict(title=title, page=index + 1) for title, index in zip(headings, positions)])
        values = [("装机", self.summary["dc_kw"], 1, 1), ("组件数", self.summary["modules"], 0, 1),
                  ("P50", self.summary["yield_kwh"]["p50"], 0, 1), ("回本", self.economics["payback_years"], 2, 1),
                  ("收益率", self.economics["irr_pct"], 1, 1), ("年减碳", self.economics["co2_t_per_year"], 0, 1),
                  ("最冷开路电压", self.electrical["voc_cold_v"], 0, 6), ("附加荷载", self.summary["load"]["added_kn_m2"], 2, 7)]
        for name, original, digits, section_index in values:
            expected = f'{Decimal(str(original)).quantize(Decimal(10) ** -digits, rounding=ROUND_HALF_UP):,.{digits}f}'
            spans = re.findall(r'title="原值 ([^"]+)"[^>]*>([^<]*)</span>', sections[section_index])
            html_ok = (str(original), expected) in [(unescape(raw), unescape(shown)) for raw, shown in spans]
            start = positions[section_index]
            end = positions[section_index + 1] if section_index + 1 < len(positions) else len(pages)
            text = "\n".join(pages[start:end]) if start >= 0 else ""
            pattern = r'\s*'.join(re.escape(character) for character in expected)
            pdf_ok = bool(re.search(r'(?<![\d,.])' + pattern + r'(?![\d,.])', text))
            self.check("proposal number: " + name, html_ok and pdf_ok, original=original, decimal_half_up=expected,
                       html_original_title=html_ok, pdf_text=pdf_ok, pdf_pages=list(range(start + 1, end + 1)))
        private = re.compile(r'account-acct|Users[\\/]|/home/', re.I)
        for kind, value in [("HTML", unescape(document)), ("PDF", "\n".join(pages))]:
            self.check(kind + " excludes private routing paths and credential patterns",
                       not private.search(value) and not SECRET.search(value) and sanitize(value) == value)
        self.report = dict(html_sha256=digest(html_path), pdf_sha256=digest(pdf_path), pdf_pages=len(pages),
                           page_text_lengths=[len(page) for page in pages])
        if preview:
            self.skip("parent report registry and project copy", "External preview selected; it is not registered to the source run")
            return
        state = self.meta.get("report", {})
        self.check("parent report state", state.get("status") == "ready" and state.get("version") == self.version,
                   status_value=state.get("status"), version=state.get("version"), generated_at=state.get("generated_at"),
                   regenerated=state.get("regenerated"))
        registered = {row["name"]: row["sha256"] for row in self.meta.get("report", {}).get("files", [])}
        for path in (html_path, pdf_path):
            actual = digest(path)
            self.check("parent report registration: " + path.name,
                       actual == registered.get(path.name) == self.meta.get("artifacts", {}).get("report/" + path.name), actual=actual)
        project = contained(ROOT / "projects", self.meta["project_id"])
        if not project.is_dir():
            self.skip("existing project report copies", "Project directory does not exist; generation must not create one")
        else:
            for path in (html_path, pdf_path):
                copy = contained(project, f"reports/v{self.version}/" + path.name)
                actual = digest(copy) if copy.is_file() else None
                self.check("existing project report copy: " + path.name, actual == digest(path), actual=actual, expected=digest(path))


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("run", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--report-only", action="store_true")
    parser.add_argument("--report-dir", type=Path)
    args = parser.parse_args(argv)
    folder, output = args.run.resolve(), args.output.resolve()
    if output.is_relative_to(folder) or any(output.is_relative_to(ROOT / name) for name in ("runs", "projects")):
        parser.error("Receipt output must not be inside any run or project directory")
    output.mkdir(parents=True, exist_ok=False)
    started = datetime.now(timezone.utc).isoformat()
    clock = time.monotonic()
    before, ignored_before = snapshot(folder)
    audit = Audit(folder)
    try:
        audit.load()
        tasks = [("proposal", lambda: audit.proposal((args.report_dir or folder / "report").resolve(), bool(args.report_dir)))]
        if args.report_only:
            audit.skip("approved DXF and XLSX", "Explicit --report-only invocation; not a full T11 acceptance")
        else:
            tasks = [("DXF", audit.dxf), ("XLSX", audit.xlsx)] + tasks
        for name, check in tasks:
            try:
                check()
            except Exception as error:
                audit.check(name + " completed", False, error=type(error).__name__ + ": " + sanitize(str(error)))
    except Exception as error:
        audit.check("archive loading", False, error=type(error).__name__ + ": " + sanitize(str(error)))
    after, ignored_after = snapshot(folder)
    audit.check("archive bytes unchanged", before == after, files_before=len(before), files_after=len(after))
    passed = not any(row["status"] == "fail" for row in audit.checks)
    receipt = dict(scope="report-only" if args.report_only else "full", passed=passed, exit_code=0 if passed else 1,
                   command=sanitize([sys.executable, *sys.argv]), run=folder.name, started_at=started,
                   ended_at=datetime.now(timezone.utc).isoformat(), elapsed_seconds=round(time.monotonic() - clock, 3),
                   sources=audit.sources, report=getattr(audit, "report", None), checks=audit.checks,
                   ignored_files_before=ignored_before, ignored_files_after=ignored_after)
    (output / "receipt.json").write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(dict(passed=passed, scope=receipt["scope"], checks=len(audit.checks),
                         failures=[row for row in audit.checks if row["status"] == "fail"]), ensure_ascii=False))
    return receipt["exit_code"]


if __name__ == "__main__":
    raise SystemExit(main())
