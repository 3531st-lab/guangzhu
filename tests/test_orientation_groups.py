"""T25 group calculations use independent sums, real rays and physical areas."""

import importlib.util
import json
import math
from pathlib import Path

import numpy as np
import pytest
import requests
from shapely.geometry import Polygon
from shapely.ops import unary_union

from guangzhu import analysis, electrical, layout, mcp_actions, orientation, project as storage
from guangzhu import rules, shading, structure, yieldsim
from guangzhu.models import Design, LayoutParams, Module, validate_project
from guangzhu.shading_geometry import build_scene, sample_points, ray_hits, sun_vector, beam_fractions
from guangzhu.solar import poa_irradiance


@pytest.fixture
def project(monkeypatch):
    def forbidden(*args, **kwargs):
        pytest.fail('T25 must be entirely offline')
    monkeypatch.setattr(requests.Session, 'request', forbidden)
    payload = json.loads(Path('data/demo/projects/changzhou-demo/project.json').read_text(encoding='utf-8'))
    payload['roof'].update(outline=[[0, 0], [100, 0], [100, 100], [0, 100]], obstacles=[], keepouts=[], parapet_height_m=0)
    payload['site'].update(design_temp_min_c=-13.45, design_temp_max_c=42.06)
    return validate_project(payload)


def panel(surface, col, x, y, *, tilt=25, azimuth=90, z=.3, row=1):
    catalog = electrical.equipment_catalog()['modules']['trina-tsm-550de19']
    width, slope = catalog['width_m'], catalog['length_m']
    angle, beta = math.radians(azimuth), math.radians(tilt)
    right = np.array([-math.cos(angle), math.sin(angle), 0]) * width
    back = np.array([-math.sin(angle) * math.cos(beta), -math.cos(angle) * math.cos(beta), math.sin(beta)]) * slope
    base = np.array([x, y, z])
    corners = [base, base + right, base + right + back, base + back]
    return Module(id=f'P{surface}-R{row:02}-1-C{col:03}', row=row, tier=1, col=col,
                  center=np.mean(corners, axis=0), corners=corners, tilt=tilt, azimuth=azimuth, orientation='portrait')


def design(modules):
    return Design(params=LayoutParams(rows_per_rack=1), modules=modules,
                  statistics=dict(gcr=.4, row_spacing_m=6, dc_kw=len(modules) * .55))


def test_group_generation_hourly_and_orientation_cache(project, monkeypatch):
    project.economics.monthly_kwh = [1000] * 12
    project.economics.shifts = 'day'
    modules = [panel(1, 1, 20, 20), panel(2, 1, 40, 20, azimuth=270), panel(3, 1, 60, 20)]
    calls = []
    def simulate(lat, lon, **kwargs):
        calls.append((kwargs['tilt'], kwargs['azimuth']))
        b = kwargs['azimuth']
        rows = {name: dict(specific_yield_kwh_kwp=value) for name, value in dict(a=b - 10, b=b, c=b + 20).items()}
        rows['b']['hourly'] = dict(utc=['a', 'b'], kwh=[b / 3, b * 2 / 3])
        return dict(results=rows, omitted={'d': 'unavailable'})
    monkeypatch.setattr(yieldsim, 'simulate_design_yields', simulate)
    shadows = dict(modules=[dict(id=m.id, total_loss_pct=loss) for m, loss in zip(modules, [10, 30, 20])])
    result, generated = analysis.generation(project, design(modules), shadows, 1.65, offline=True)
    expected = sum(.55 * b * (1 - loss / 100) for b, loss in zip([90, 270, 90], [10, 30, 20]))
    assert calls == [(25, 90), (25, 270)]
    assert result['p50'] == pytest.approx(expected)
    assert result['low'] == pytest.approx(expected - 10 * .55 * 2.4)
    assert result['high'] == pytest.approx(expected + 20 * .55 * 2.4)
    assert sum(result['hourly']['kwh']) == pytest.approx(expected)
    assert result['per_kw'] == pytest.approx(expected / 1.65)
    assert len(generated['groups']) == 3
    assert 'hourly' not in json.dumps(yieldsim.without_hourly(generated))
    one = design(modules[:1])
    one_shadow = dict(summary=dict(avg_loss_pct=10), modules=shadows['modules'][:1])
    value, _ = analysis.generation(project, one, one_shadow, .55, offline=True)
    assert calls[-1] == (25, 90)
    assert value['p50'] == pytest.approx(.55 * 90 * .9)


def test_real_shading_uses_own_poa_and_other_group_blocker(project):
    target = panel(1, 1, 20, 20, azimuth=90)
    blocker = panel(2, 1, 20, 20, azimuth=270, z=2)
    mixed = design([target, blocker])
    site = shading.SiteShading(project, offline=True)
    result = shading.analyze_design(project, mixed, _site=site)
    scene, points = build_scene(project, mixed), sample_points(mixed)
    owners = np.broadcast_to(np.arange(2)[:, None, None], points.shape[:-1])
    expected = []
    for index, module in enumerate(mixed.modules):
        poa = poa_irradiance(site.weather, site.sun, module.tilt, module.azimuth, project.site.albedo)
        weights = np.zeros((12, 24))
        np.add.at(weights, (site.weather.index.month - 1, site.weather.index.hour), poa.poa_direct.to_numpy())
        lost = 0
        for instant, position in site.positions.iterrows():
            if position.elevation > 0:
                hit = ray_hits(scene, points, sun_vector(position.elevation, position.azimuth), owners)
                fraction, _ = beam_fractions(hit, [m.orientation for m in mixed.modules])
                lost += weights[instant.month - 1, instant.hour] * fraction[index]
        expected.append(100 * lost / poa.poa_global.sum())
    assert [m['beam_loss_pct'] for m in result['modules']] == pytest.approx(expected)
    isolated = shading.analyze_design(project, design([target]), _site=site)
    assert result['modules'][0]['beam_loss_pct'] > isolated['modules'][0]['beam_loss_pct']
    reused = shading.RemovalShading(project, mixed, site=site)
    assert reused.analyze(mixed)['modules'] == result['modules']
    reduced = mixed.model_copy(update=dict(modules=[target]))
    assert reused.analyze(reduced)['modules'] == shading.analyze_design(project, reduced, _site=site)['modules']


@pytest.mark.parametrize('counts', [(48, 48), (13, 50), (47, 49), (16, 16, 16, 16)])
def test_strings_and_mppt_reservations_never_mix_groups(project, counts):
    modules = [panel(g + 1, i + 1, 10 + g * 15, 10 + i, azimuth=90 if g % 2 else 270)
               for g, count in enumerate(counts) for i in range(count)]
    result = electrical.calculate(project, design(modules), offline=True)
    identities = {m.id: orientation.key(m) for m in modules}
    assigned, ports = [], {}
    for string in result['strings']:
        groups = {identities[identifier] for identifier in string['module_ids']}
        assert len(groups) == 1
        ports.setdefault((string['inverter'], string['mppt']), set()).add((groups.pop(), string['module_count']))
        assigned += string['module_ids']
    assert all(len(tokens) == 1 for tokens in ports.values())
    assert len(assigned) == len(set(assigned))
    assert set(assigned).isdisjoint(result['unused_modules'])
    assert set(assigned) | set(result['unused_modules']) == set(identities)
    if counts == (48, 48):
        assert result['inverter_count'] == 1
        assert not result['unused_modules']


def test_load_sums_physical_projections_and_preserves_surface_ids(project):
    modules = [panel(1, 1, 20, 20, tilt=10), panel(2, 1, 40, 20, tilt=30, azimuth=270)]
    result = structure.calculate(project, design(modules))
    expected_area = unary_union([Polygon([p[:2] for p in m.corners]) for m in modules]).area
    assert result['rack_projection_area_m2'] == pytest.approx(expected_area)
    assert result['added_kn_m2'] == pytest.approx(result['mass_kg'] * 9.81 / 1000 / expected_area)


def test_g02_checks_each_group_independently(project):
    groups = [layout.generate_layout(project, dict(azimuth_deg=azimuth, rows_per_rack=1)) for azimuth in [90, 270]]
    modules = []
    for number, source in enumerate(groups, 1):
        selected = [m.model_copy(update=dict(id=f'P{number}-' + m.id)) for m in source.modules if m.row <= 2 and m.col == 1]
        modules.extend(selected)
    mixed = groups[0].model_copy(update=dict(modules=modules))
    electric = electrical.calculate(project, mixed, offline=True)
    shade = dict(summary=dict(avg_loss_pct=0, modules_over_threshold=0), modules=[])
    result = rules.evaluate(project, mixed, shade, electric, structure.calculate(project, mixed))
    checks = [row for row in result['items'] if row['id'] == 'G02']
    assert len(checks) == 2
    assert [row['group'][2] for row in checks] == [90, 270]
    for (_, view), check in zip(orientation.views(mixed), checks):
        length, gaps, _ = rules.row_geometry(project, view)
        expected = layout.row_spacing(project.site.lat, project.site.lon, view.params.tilt_deg, length, view.params.azimuth_deg)
        assert check['limit'] == expected['row_spacing_m']
        assert check['value'] == min(gaps)


def test_prefixed_modules_delete_and_compress(project, monkeypatch, tmp_path):
    monkeypatch.setattr(storage, 'PROJECTS_DIR', tmp_path)
    source = Path('data/demo/projects/changzhou-demo')
    instance = storage.init(source, 't25-prefix')
    record = design([panel(2, i, 20, 10 + i) for i in range(1, 10)])
    storage.save_design(instance.id, record)
    operation = dict(op='remove_modules', module_ids=[m.id for m in record.modules[:8]])
    assert analysis._compact_removal(operation)['module_id_ranges'] == ['P2-R01-1-C001..P2-R01-1-C008']
    response = storage.edit_design(instance.id, 1, [operation], '测试删除', 'rule')
    edited = storage.load_design(instance.id, response['version'])
    assert len(edited.modules) == 1
    assert edited.modules[0].id == 'P2-R01-1-C009'


def test_baseline_field_diff():
    spec = importlib.util.spec_from_file_location('baseline', Path('scripts/baseline.py'))
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    assert module.differences(dict(a=[1, 2]), dict(a=[1, 3])) == ['$.a[1]']
    assert module.differences(dict(a=1), dict(a=1, b=2)) == ['$.b']


def test_manifest_detects_tampering_and_missing_files(tmp_path, monkeypatch):
    import hashlib
    spec = importlib.util.spec_from_file_location('baseline', Path('scripts/baseline.py'))
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    payload = b'{"a": 1}'
    (tmp_path / 'roof-summary.json').write_bytes(payload)
    manifest = tmp_path / 'manifest.json'
    manifest.write_text(json.dumps({'roof-summary.json': hashlib.sha256(payload).hexdigest()}), encoding='utf-8')
    assert module.reference_manifest(manifest)[0] == tmp_path
    monkeypatch.setattr('sys.argv', ['baseline.py', '--output', str(tmp_path / 'new'), '--compare', str(manifest)])
    with pytest.raises(ValueError, match='manifest entries differ'):
        module.main()
    (tmp_path / 'roof-summary.json').write_bytes(b'{"a": 2}')
    with pytest.raises(ValueError, match='checksum mismatch'):
        module.reference_manifest(manifest)
    (tmp_path / 'roof-summary.json').unlink()
    with pytest.raises(FileNotFoundError):
        module.reference_manifest(manifest)


def test_group_geometry_does_not_reuse_global_rack_length_or_missing_tier(project):
    from guangzhu.models import RemovedModule
    modules, removed, references = [], [], []
    for surface, azimuth, mode, tiers in [(1, 90, 'portrait', 2), (2, 270, 'landscape', 1)]:
        original = layout.generate_layout(project, dict(azimuth_deg=azimuth, orientation=mode, rows_per_rack=tiers))
        selected = [m.model_copy(update=dict(id=f'P{surface}-' + m.id)) for m in original.modules if m.row <= 2 and m.col == 1]
        references.append(original.model_copy(update=dict(modules=selected)))
        for m in selected:
            if surface == 1 and m.row == 2 and m.tier == 1:
                removed.append(RemovedModule(**m.model_dump(), reason='测试缺层'))
            else:
                modules.append(m)
    mixed = references[0].model_copy(update=dict(modules=modules, removed_modules=removed))
    grouped = list(orientation.views(mixed))
    assert [view.params.rows_per_rack for _, view in grouped] == [2, 1]
    lengths = []
    for (_, view), original in zip(grouped, references):
        length, gaps, _ = rules.row_geometry(project, view)
        old_length, old_gaps, _ = rules.row_geometry(project, original)
        assert length == pytest.approx(old_length, abs=.002)
        assert gaps == pytest.approx(old_gaps, abs=.002)
        lengths.append(length)
    assert lengths[0] > 3 * lengths[1]
    expected_area = sum(structure.calculate(project, original)['rack_projection_area_m2'] for original in references)
    assert structure.calculate(project, mixed)['rack_projection_area_m2'] == pytest.approx(expected_area)


def test_g05_uses_each_module_group_tilt(project):
    project.requirements.structured.tilt_deg_range = (20, 30)
    mixed = design([panel(1, 1, 20, 20, tilt=10), panel(2, 1, 40, 20, tilt=25, azimuth=270)])
    electric = electrical.calculate(project, mixed, offline=True)
    shade = dict(summary=dict(avg_loss_pct=0, modules_over_threshold=0), modules=[])
    checked = rules.evaluate(project, mixed, shade, electric, structure.calculate(project, mixed))
    assert [(row['value'], row['status']) for row in checked['items'] if row['id'] == 'G05'] == [(10, 'fail'), (25, 'pass')]


def test_crosscheck_pagination_preserves_all_groups_and_full_totals(project, monkeypatch, tmp_path):
    from guangzhu import mcp_server
    from guangzhu.models import GzError
    mixed = design([panel(i + 1, 1, 20, 20, azimuth=90 if i % 2 else 270) for i in range(45)])
    monkeypatch.setattr(storage, 'load', lambda _: project)
    monkeypatch.setattr(storage, 'load_design', lambda *_: mixed)
    monkeypatch.setattr(storage, 'PROJECTS_DIR', tmp_path)
    calls = []
    def simulate(lat, lon, **kwargs):
        calls.append(kwargs['azimuth'])
        return dict(results={name: dict(specific_yield_kwh_kwp=kwargs['azimuth'] + i)
                             for i, name in enumerate(['a', 'b', 'c', 'b_tmy', 'd'])}, omitted={})
    monkeypatch.setattr(yieldsim, 'simulate_design_yields', simulate)
    gathered, offset, sizes = [], 0, []
    while offset is not None:
        response = mcp_server.call('gz_crosscheck_yield', dict(project_id=project.id, version=1, group_offset=offset))
        assert not response.is_error
        result = json.loads(response.content[0].text)
        assert result['group_offset'] == offset and result['total_groups'] == 45
        assert result['groups']
        # JobQueue prepends bounded job metadata to the serialized result.
        assert len(response.content[0].text.encode('utf-8')) + mcp_server.WAIT_METADATA_BYTES <= 2048
        wrapped = mcp_server.JobQueue.response(dict(result, tool='gz_crosscheck_yield', job_id='0' * 32))
        sizes.append(len(wrapped.content[0].text.encode('utf-8')))
        assert sizes[-1] <= 2048
        gathered.extend(result['groups'])
        assert result['total_kwh']['b'] == pytest.approx(round(sum(.55 * (m.azimuth + 1) for m in mixed.modules), 4))
        offset = result['next_group_offset']
    assert [row['group'][0] for row in gathered] == list(range(1, 46))
    assert len(calls) == len(sizes) * 2
    assert len(json.loads(storage._path(project.id, 'analysis/v1/crosscheck.json').read_text(encoding='utf-8'))['groups']) == 45
    for invalid in [-1, 45, True]:
        with pytest.raises(GzError, match='朝向组'):
            mcp_actions._crosscheck_yield(project.id, 1, group_offset=invalid)
    print('T25 wrapped pagination byte sizes:', sizes)


def test_grouped_analysis_and_materials_use_actual_models(project, monkeypatch, tmp_path):
    from openpyxl import load_workbook
    from guangzhu.report.deliverables import write_deliverables
    monkeypatch.setattr(storage, 'PROJECTS_DIR', tmp_path)
    instance = storage.init(Path('data/demo/projects/changzhou-demo'), 't25-real-groups')
    modules = [panel(g + 1, i + 1, 20 + g * 25, 10 + i * 1.2, tilt=15 + 10 * g, azimuth=90 + 180 * g)
               for g in range(2) for i in range(16)]
    mixed = design(modules)
    storage.save_design(instance.id, mixed)
    result = analysis.analyze(instance.id, 1, offline=True)
    assert result['ok'] and result['modules'] == 32
    assert len(analysis.serialized(result).encode('utf-8')) <= 2048
    cash = json.loads(storage._path(instance.id, 'analysis/v1/economics.json').read_text(encoding='utf-8'))
    assert len(cash['generation']['groups']) == 2
    assert '"hourly":' not in json.dumps(cash)
    crosscheck = mcp_actions._crosscheck_yield(instance.id, 1, offline=True)
    assert crosscheck['next_group_offset'] is None and len(crosscheck['groups']) == 2
    expected = sum(g['kw'] * g['per_kw']['b'] for g in crosscheck['groups'])
    assert crosscheck['total_kwh']['b'] == pytest.approx(expected, abs=.001)
    directory = tmp_path / 'exports'
    directory.mkdir()
    write_deliverables(instance.id, 1, directory)
    book = load_workbook(directory / 'bom-v1.xlsx')
    assert book['朝向组'].max_row == 3
    assert book['朝向组']['D2'].value == 16 and book['朝向组']['D3'].value == 16
    assert len(list(book['组串表'].values)) == 3


def test_api_uses_complete_persisted_crosscheck_instead_of_one_page():
    from guangzhu import api
    def node(count):
        return dict(kind='tool', name='gz_crosscheck_yield', result_json=dict(
            ok=True, version=1, total_groups=3, groups=[dict(group=[i, 25, 90]) for i in range(count)]))
    assert api.tool_result([node(1), node(2)], 'gz_crosscheck_yield', 1) is None
    assert len(api.tool_result([node(3)], 'gz_crosscheck_yield', 1)['groups']) == 3


def test_single_remaining_group_rebuilds_geometry_but_legacy_removal_keeps_old_view(project):
    from guangzhu.models import RemovedModule
    old = layout.generate_layout(project, dict(rows_per_rack=2))
    target = old.modules[0]
    removed = RemovedModule(**old.modules[1].model_dump(), reason='普通删板')
    ordinary = old.model_copy(update=dict(modules=[target], removed_modules=[removed]))
    assert next(orientation.views(ordinary))[1] is ordinary
    other = RemovedModule(**panel(2, 1, 80, 80, tilt=10).model_dump(), reason='另一朝向全部删除')
    remaining = old.model_copy(update=dict(modules=[target], removed_modules=[other]))
    view = next(orientation.views(remaining))[1]
    assert view is not remaining and view.params.rows_per_rack == 1
    expected = electrical.equipment_catalog()['modules'][project.equipment.module_id]['length_m']
    assert rules.row_geometry(project, view)[0] == pytest.approx(expected, abs=.002)
    prefixed = old.model_copy(update=dict(modules=[target.model_copy(update=dict(id='P1-' + target.id))]))
    prefixed_view = next(orientation.views(prefixed))[1]
    assert prefixed_view.params.rows_per_rack == 1
    assert rules.row_geometry(project, prefixed_view)[0] == pytest.approx(expected, abs=.002)
    old_load = structure.calculate(project, ordinary)
    new_load = structure.calculate(project, prefixed)
    assert old_load['rack_projection_area_m2'] > 2 * new_load['rack_projection_area_m2']


@pytest.mark.parametrize('mix_inside_string', [False, True])
def test_rules_cannot_overwrite_cross_group_e09_failure_with_same_length_pass(project, mix_inside_string):
    import copy
    modules = [panel(g + 1, i + 1, 20 + g * 20, 10 + i, azimuth=90 + 180 * g) for g in range(2) for i in range(16)]
    mixed = design(modules)
    electric = copy.deepcopy(electrical.calculate(project, mixed, offline=True))
    first, second = electric['strings']
    assert len(first['module_ids']) == len(second['module_ids']) == 16
    if mix_inside_string:
        first['module_ids'][0], second['module_ids'][0] = second['module_ids'][0], first['module_ids'][0]
    else:
        second.update(inverter=first['inverter'], mppt=first['mppt'])
    electric['checks']['E09']['status'] = 'fail'
    shade = dict(summary=dict(avg_loss_pct=0, modules_over_threshold=0), modules=[])
    checked = rules.evaluate(project, mixed, shade, electric, structure.calculate(project, mixed))
    check = next(row for row in checked['items'] if row['id'] == 'E09')
    assert check['status'] == 'fail'
    assert '朝向组数' in check['evidence']


def test_single_prefixed_north_group_uses_group_spacing_domain(project):
    mixed = design([panel(1, 1, 20, 20, azimuth=0)])
    electric = electrical.calculate(project, mixed, offline=True)
    shade = dict(summary=dict(avg_loss_pct=0, modules_over_threshold=0), modules=[])
    checked = rules.evaluate(project, mixed, shade, electric, structure.calculate(project, mixed))
    assert next(row for row in checked['items'] if row['id'] == 'G02')['status'] == 'pass'


def test_deleting_all_prefixed_modules_preserves_empty_production_analysis(project, monkeypatch, tmp_path):
    monkeypatch.setattr(storage, 'PROJECTS_DIR', tmp_path)
    instance = storage.init(Path('data/demo/projects/changzhou-demo'), 't25-empty-groups')
    mixed = design([panel(1, 1, 20, 20), panel(2, 1, 40, 20, azimuth=270)])
    storage.save_design(instance.id, mixed)
    edit = storage.edit_design(instance.id, 1,
                               [dict(op='remove_modules', module_ids=[m.id for m in mixed.modules])], '全部删除', 'rule')
    empty = storage.load_design(instance.id, edit['version'])
    assert empty.modules == [] and len(empty.removed_modules) == 2
    summary = analysis.analyze(instance.id, empty.version, offline=True)
    assert summary['ok'] and summary['modules'] == 0 and summary['dc_kw'] == 0
    assert [summary['yield_kwh'][name] for name in ('p50', 'low', 'high')] == [0, 0, 0]
    assert summary['shading']['avg_loss_pct'] == summary['load']['added_kn_m2'] == 0
    payloads = {}
    for name in ('shading', 'structure', 'electrical', 'rules', 'economics'):
        path = storage._path(instance.id, f'analysis/v{empty.version}/{name}.json')
        payloads[name] = json.loads(path.read_text(encoding='utf-8'))
    assert payloads['shading']['modules'] == []
    assert payloads['structure']['mass_kg'] == payloads['structure']['rack_projection_area_m2'] == 0
    assert payloads['electrical']['strings'] == payloads['electrical']['unused_modules'] == []
    assert next(row for row in payloads['rules']['items'] if row['id'] == 'E09')['status'] == 'pass'
    assert payloads['economics']['analysis_metadata']['yield_kwh']['p50'] == 0
    legacy = empty.model_copy(update=dict(removed_modules=[]))
    assert next(orientation.views(empty))[1] is empty
    assert structure.calculate(instance, empty) == structure.calculate(instance, legacy)
    assert mcp_actions._crosscheck_yield(instance.id, empty.version, offline=True)['ok']
