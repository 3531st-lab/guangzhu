"""The acceptance reader can verify actual binary formats and catches tampering."""
import hashlib
import json
import shutil

import pytest

from test_report_deliverables import exported
from validate_t11_run import Audit, points_match, snapshot


def actual_audit(root, project_id):
    base = root / "projects" / project_id
    audit = Audit(root)
    audit.capture = root
    audit.export = base / "exports/v3"
    audit.version = 3
    read = lambda relative: json.loads((base / relative).read_text(encoding="utf-8"))
    audit.project = read("project.json")
    audit.design = read("designs/v3.json")
    audit.electrical = read("analysis/v3/electrical.json")
    audit.economics = read("analysis/v3/economics.json")
    audit.manifest = read("exports/v3/manifest.json")
    audit.capture_meta = dict(artifacts={path.relative_to(root).as_posix(): hashlib.sha256(path.read_bytes()).hexdigest()
                                      for path in audit.export.iterdir()})
    return audit


def test_acceptance_binary_checks_use_actual_offline_export_fixture(exported, monkeypatch):
    root, project_id, _ = exported
    audit = actual_audit(root, project_id)
    audit.dxf()
    audit.xlsx()
    assert not [row for row in audit.checks if row["status"] == "fail"]
    assert next(row for row in audit.checks if row["name"] == "XLSX 26 cashflow rows equal archived values")["rows"] == 26
    import openpyxl
    book = openpyxl.load_workbook(audit.export / 'bom-v3.xlsx', data_only=False)
    inverter_row = next(row for row in book['材料清单'] if row[0].value == '逆变器')
    audit.electrical['inverters'] = [dict(model='大机', count=1), dict(model='小机', count=1)]
    with monkeypatch.context() as patch:
        patch.setattr(openpyxl, 'load_workbook', lambda *args, **kwargs: book)
        inverter_row[2].value = audit.electrical['inverter_model']
        audit.checks.clear()
        audit.xlsx()
        counts = next(row for row in audit.checks if row['name'] == 'XLSX counts capacity and inverter model')
        assert counts['status'] == 'fail', 'mixed result must not accept a bare main model'
        inverter_row[2].value = '大机 × 1 + 小机 × 1'
        audit.checks.clear()
        audit.xlsx()
        counts = next(row for row in audit.checks if row['name'] == 'XLSX counts capacity and inverter model')
        assert counts['status'] == 'pass'
    audit.manifest["files"]["layout-v3.dxf"] = "0" * 64
    audit.binary("dxf")
    assert audit.checks[-1]["status"] == "fail"


@pytest.mark.parametrize('variant', ['legacy', 'current', '材料清单', '组串表', '经济测算', '假设与来源', 'cashflow', 'formula'])
def test_xlsx_core_schema_and_tamper_checks_use_real_workbooks(exported, tmp_path, variant):
    import openpyxl
    root, project_id, _ = exported
    shutil.copytree(root / 'projects', tmp_path / 'projects')
    audit = actual_audit(tmp_path, project_id)
    path = audit.export / 'bom-v3.xlsx'
    book = openpyxl.load_workbook(path)
    core = {'材料清单', '组串表', '经济测算', '假设与来源'}
    if variant == 'legacy':
        for name in list(book.sheetnames):
            if name not in core:
                del book[name]
        assert len(book.sheetnames) == 4
    elif variant in core:
        del book[variant]
    elif variant == 'cashflow':
        book['经济测算']['G11'] = 123456789
    elif variant == 'formula':
        book['造价明细']['A2'] = '=1+1'
    else:
        assert len(book.sheetnames) == 11 and {'风荷载', '串联布线'} <= set(book.sheetnames)
    book.save(path)
    book.close()
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    audit.manifest['files'][path.name] = digest
    audit.capture_meta['artifacts'][path.relative_to(tmp_path).as_posix()] = digest
    audit.xlsx()
    failed = [row['name'] for row in audit.checks if row['status'] == 'fail']
    if variant in core:
        assert failed == ['XLSX four core sheets']
        assert audit.checks[-1]['missing'] == [variant]
    elif variant == 'cashflow':
        assert failed == ['XLSX 26 cashflow rows equal archived values']
    elif variant == 'formula':
        assert failed == ['XLSX no executable formulas']
    else:
        assert not failed


def test_geometry_tolerance_checks_every_coordinate_and_vertex_count():
    assert points_match([(1, 2), (3, 4)], [(1.0009, 2), (3, 4)])
    assert not points_match([(1, 2), (3, 4)], [(1.0011, 2), (3, 4)])
    assert not points_match([(1, 2)], [(1, 2), (3, 4)])


@pytest.mark.parametrize("change", [None, "dc_length", "ac_length", "dc_section", "ac_section", "unit",
                                    "missing", "duplicate", "extra", "invented_power", "missing_note", "boolean"])
def test_calculated_cable_audit_checks_every_section_and_rejects_tampering(tmp_path, change):
    audit = Audit(tmp_path)
    audit.electrical = dict(cables=dict(summary=[dict(kind="dc", section_mm2=4, length_m=12.5),
                                               dict(kind="dc", section_mm2=6, length_m=27),
                                               dict(kind="ac", section_mm2=10, length_m=4.75)]))
    rows = [["直流电缆", None, "4 mm²", None, 12.5, "m", None, "示意"],
            ["直流电缆", None, "6 mm²", None, 27, "m", None, "示意"],
            ["交流电缆", None, "4×10 mm²", None, 4.75, "m", None, "示意"]]
    edits = {"dc_length": (0, 4, 13), "ac_length": (2, 4, 19), "dc_section": (0, 2, "10 mm²"),
             "ac_section": (2, 2, "10 mm²"), "unit": (0, 5, "km"), "invented_power": (0, 6, 5),
             "missing_note": (0, 7, None), "boolean": (0, 4, True)}
    if change in edits:
        row, column, value = edits[change]
        rows[row][column] = value
    elif change == "missing":
        rows.pop(0)
    elif change == "duplicate":
        rows[0] = rows[1].copy()
    elif change == "extra":
        rows.append(rows[0].copy())
    audit.xlsx_cables([tuple(row) for row in reversed(rows)])
    assert audit.checks[-1]["status"] == ("pass" if change is None else "fail")


@pytest.mark.parametrize("column", [None, 1, 2, 3, 4, 5, 6, "extra_ac"])
def test_legacy_cable_audit_preserves_uncalculated_contract(tmp_path, column):
    audit = Audit(tmp_path)
    audit.electrical = {}
    row = ["直流电缆", None, None, None, None, None, None, "本期未计算"]
    if isinstance(column, int):
        row[column] = 10
    rows = [tuple(row)]
    if column == "extra_ac":
        rows.append(("交流电缆", None, "4×6 mm²", None, 100, "m", None, "编造"))
    audit.xlsx_cables(rows)
    assert audit.checks[-1]["status"] == ("pass" if column is None else "fail")


def test_snapshot_only_ignores_unregistered_office_lock_without_reading_it(tmp_path, monkeypatch):
    from pathlib import Path
    import pytest
    metadata = tmp_path / "run.json"
    metadata.write_text('{"artifacts":{}}', encoding="utf-8")
    lock = tmp_path / "~$bom-v12.xlsx"
    lock.write_bytes(b"test lock")
    original = Path.read_bytes
    monkeypatch.setattr(Path, "read_bytes", lambda path: (_ for _ in ()).throw(PermissionError("locked")) if path == lock else original(path))
    hashes, ignored = snapshot(tmp_path)
    assert list(hashes) == ["run.json"]
    assert [row["path"] for row in ignored] == [lock.name]
    metadata.write_text(json.dumps(dict(artifacts={lock.name: "registered"})), encoding="utf-8")
    with pytest.raises(PermissionError):
        snapshot(tmp_path)
