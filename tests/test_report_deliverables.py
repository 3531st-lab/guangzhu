"""Read exported engineering files back with the actual format libraries."""

import hashlib
import json
import shutil
from pathlib import Path

import ezdxf
import openpyxl
import pytest
import requests

from guangzhu import analysis, electrical, mcp_actions, project as storage
from guangzhu.config import ROOT
from guangzhu.models import GzError
from guangzhu.report.deliverables import write_deliverables
from scripts import agh_export_run as exporter


def load(path):
    return json.loads(path.read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def exported(tmp_path_factory):
    root = tmp_path_factory.mktemp("deliverables")
    with pytest.MonkeyPatch.context() as patch:
        patch.setattr(storage, "PROJECTS_DIR", root / "projects")
        patch.setattr(requests.Session, "request", lambda *a, **k: pytest.fail("offline export contacted network"))
        project = storage.init(ROOT / "data/demo/projects/changzhou-demo")
        project.roof.inverter_positions = [(30, 18)]  # Explicit closer placement in this disposable export fixture.
        project.roof.load_capacity_kn_m2 = 2  # Export-format fixture; original 0.5 now correctly blocks on wind L01.
        storage.save(project)
        result = mcp_actions._export(project.id, 3, offline=True)
        electric = load(storage._path(project.id, 'analysis/v3/electrical.json'))
        assert electric['checks']['E10']['status'] == electric['checks']['E11']['status'] == 'pass'
        yield root, project.id, result


def test_dxf_geometry_layers_units_labels_and_bbox_match_design(exported):
    root, project_id, result = exported
    project = load(root / "projects" / project_id / "project.json")
    design = load(root / "projects" / project_id / "designs/v3.json")
    electric = load(root / "projects" / project_id / "analysis/v3/electrical.json")
    drawing = ezdxf.readfile(root / "projects" / project_id / "exports/v3/layout-v3.dxf")
    space = drawing.modelspace()
    assert drawing.units == drawing.header["$INSUNITS"] == 6
    assert {"ROOF", "OBSTACLE", "KEEPOUT", "PV", "STRING", "TEXT"} <= {layer.dxf.name for layer in drawing.layers}
    assert drawing.layers.get("KEEPOUT").dxf.linetype == "DASHED"
    polygons = list(space.query('LWPOLYLINE[layer=="PV"]'))
    assert len(polygons) == len(design["modules"])
    for actual, expected in zip(polygons, design["modules"]):
        assert actual.closed
        for point, corner in zip(actual.get_points("xy"), expected["corners"]):
            assert point == pytest.approx(corner[:2], abs=.001)
    points = [p for shape in polygons for p in shape.get_points("xy")]
    original = [p for module in design["modules"] for p in module["corners"]]
    for axis in (0, 1):
        assert min(p[axis] for p in points) == pytest.approx(min(p[axis] for p in original), abs=.001)
        assert max(p[axis] for p in points) == pytest.approx(max(p[axis] for p in original), abs=.001)
    roof = list(space.query('LWPOLYLINE[layer=="ROOF"]'))
    assert len(roof) == 1 and roof[0].closed
    assert list(roof[0].get_points("xy")) == [tuple(p) for p in project["roof"]["outline"]]
    boxes = [o for o in project["roof"]["obstacles"] if o["kind"] != "cylinder"]
    circles = [o for o in project["roof"]["obstacles"] if o["kind"] == "cylinder"]
    assert len(space.query('LWPOLYLINE[layer=="OBSTACLE"]')) == len(boxes)
    for actual, expected in zip(space.query('CIRCLE[layer=="OBSTACLE"]'), circles):
        assert tuple(actual.dxf.center)[:2] == tuple(expected["center"])
        assert actual.dxf.radius == expected["radius_m"]
    assert len(space.query('LWPOLYLINE[layer=="KEEPOUT"]')) == len(project["roof"]["keepouts"]) + len(design["added_keepouts"])
    labels = list(space.query('TEXT[layer=="STRING"]'))
    panels = {p["id"]: p for p in design["modules"]}
    assert [label.dxf.text for label in labels] == [s["id"] for s in electric["strings"]]
    for label, string in zip(labels, electric["strings"]):
        assert tuple(label.dxf.insert)[:2] == tuple(panels[string["module_ids"][0]]["corners"][0][:2])
    title = " ".join(row.dxf.text for row in space.query('TEXT[layer=="TEXT"]'))
    assert project["name"] in title and "v3" in title and "+08:00" in title


def test_xlsx_original_counts_currents_mass_cashflows_and_sources(exported):
    root, project_id, _ = exported
    base = root / "projects" / project_id
    electric = load(base / "analysis/v3/electrical.json")
    economy = load(base / "analysis/v3/economics.json")
    structure = load(base / "analysis/v3/structure.json")
    design = load(base / "designs/v3.json")
    book = openpyxl.load_workbook(base / "exports/v3/bom-v3.xlsx", data_only=False)
    assert book.sheetnames == ["材料清单", "组串表", "串联布线", "经济测算", "假设与来源", '风荷载', "造价明细", "自投年度账",
                              "合同能源管理年度账", "屋顶租赁年度账", "分时与收入"]
    materials = {row[0]: row for row in book["材料清单"].iter_rows(min_row=2, values_only=True)}
    assert materials["组件"][4] == materials["压载支架"][4] == len(design["modules"])
    assert materials["组件"][6] == design["statistics"]["dc_kw"]
    assert materials["逆变器"][4] == electric["inverter_count"]
    assert materials["组件、支架及压块总质量"][6] == pytest.approx(structure["mass_kg"], rel=1e-12)
    assert "含组件" in materials["组件、支架及压块总质量"][7]
    assert materials["压载支架"][7] == "示意"
    assert materials["压载支架"][2] == "压载斜支架"
    rack_cost = next(row['subtotal_cny'] for row in economy['costs']['items'] if row['key'] == 'ballast_tilt')
    assert materials['压载支架'][9] == pytest.approx(rack_cost)
    assert materials['压载支架'][8] * materials['压载支架'][4] == pytest.approx(rack_cost)
    rack_rows = [row for name, row in materials.items() if '支架' in name]
    assert sum(row[9] for row in rack_rows if isinstance(row[9], (int, float))) == pytest.approx(rack_cost)
    assert materials['组件、支架及压块总质量'][8:10] == ('不适用', '不适用')
    assert materials['组件、支架及压块总质量'][10] == '工程统计行，不重复计采购费用'
    assert load(base / "project.json")["equipment"]["racking"] == "ballast_tilt"
    assert "示意" in materials["组件、支架及压块总质量"][7]
    assert materials["组串"][4] == electric["string_count"]
    assert materials["每串块数"][4] == electric["modules_per_string"]
    cable_rows = [r for r in book['材料清单'].iter_rows(min_row=2, values_only=True) if r[0] in ['直流电缆', '交流电缆']]
    assert [r[4] for r in cable_rows] == pytest.approx([r['length_m'] for r in electric['cables']['summary']])
    assert all(r[5] == 'm' and '示意' in r[7] for r in cable_rows)
    currents = {(r["inverter"], r["mppt"]): r for r in electric["mppt_currents"]}
    assert "MPPT 合计" in book["组串表"].cell(1, 7).value
    for row, string in zip(book["组串表"].iter_rows(min_row=2, values_only=True), electric["strings"]):
        current = currents[string["inverter"], string["mppt"]]
        expected = [string["id"], len(string["module_ids"]), string["inverter"], string["mppt"],
                    electric["voc_cold_v"] * len(string['module_ids']) / electric['modules_per_string'],
                    electric["vmp_hot_v"] * len(string['module_ids']) / electric.get('modules_per_string_short', electric['modules_per_string']),
                    current["working_a"], current["short_a"]]
        assert row[0] == expected[0] and row[1:] == pytest.approx(expected[1:], rel=1e-12)
    keys = ["year", "generation_kwh", "self_use_kwh", "export_kwh", "income_cny", "om_cny", "net_cny",
            "cumulative_cny", "discounted_cny", "cumulative_discounted_cny"]
    rows = list(book["经济测算"].iter_rows(min_row=10, values_only=True))
    assert len(rows) == 26 and [r[0] for r in rows] == list(range(26))
    for row, original in zip(rows, economy["cashflows"]):
        assert row == pytest.approx([original[key] for key in keys], rel=1e-12)
    for row, (key, source) in zip(book["假设与来源"].iter_rows(min_row=2, values_only=True), economy["source"].items()):
        assert row[1] == pytest.approx(economy["params"][key], rel=1e-12)
        assert (row[0], *row[2:]) == (source["中文名"], source["单位"], source["状态"],
                                     source["来源标题"] or None, source["链接"] or None, source["访问日期"] or None)


def test_export_preserves_four_json_bytes_and_small_hash_manifest(exported):
    root, project_id, result = exported
    base = root / "projects" / project_id
    for name, path in {"design": "designs/v3.json", "summary": "analysis/v3/summary.json",
                       "rules": "analysis/v3/rules.json", "economics": "analysis/v3/economics.json"}.items():
        assert (base / f"exports/v3/{name}.json").read_bytes() == (base / path).read_bytes()
    assert len(result["files"]) == 6
    assert len(json.dumps(result, ensure_ascii=False).encode("utf-8")) <= 2048
    for name, digest in result["files"].items():
        assert hashlib.sha256((base / "exports/v3" / name).read_bytes()).hexdigest() == digest


@pytest.mark.parametrize("legacy", [False, True])
def test_repeated_export_does_not_rewrite_approved_bundle(exported, tmp_path, monkeypatch, legacy):
    root, project_id, _ = exported
    shutil.copytree(root / "projects", tmp_path / "projects")
    monkeypatch.setattr(storage, "PROJECTS_DIR", tmp_path / "projects")
    monkeypatch.setattr(analysis, "analyze", lambda *a, **k: None)
    base = tmp_path / "projects" / project_id / "exports/v3"
    if legacy:
        manifest = load(base / "manifest.json")
        for name in ("layout-v3.dxf", "bom-v3.xlsx"):
            del manifest["files"][name]
            (base / name).unlink()
        (base / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    else:
        book = openpyxl.load_workbook(base / "bom-v3.xlsx")
        book["材料清单"]["C4"] = "ballast_tilt"
        book.save(base / "bom-v3.xlsx")
        manifest = load(base / "manifest.json")
        manifest["files"]["bom-v3.xlsx"] = hashlib.sha256((base / "bom-v3.xlsx").read_bytes()).hexdigest()
        (base / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    before = {p.name: p.read_bytes() for p in base.iterdir()}
    result = mcp_actions._export(project_id, 3, offline=True)
    assert {p.name: p.read_bytes() for p in base.iterdir()} == before
    assert len(result["files"]) == (4 if legacy else 6)
    (base / "design.json").write_bytes(b"tampered")
    with pytest.raises(GzError, match="原导出包"):
        mcp_actions._export(project_id, 3, offline=True)
    assert (base / "design.json").read_bytes() == b"tampered"


def test_xlsx_source_text_cannot_become_a_formula(exported, tmp_path, monkeypatch):
    _, project_id, _ = exported
    catalog = electrical.equipment_catalog()
    catalog["modules"]["trina-tsm-550de19"]["manufacturer"] = "=1+1"
    monkeypatch.setattr(electrical, "equipment_catalog", lambda: catalog)
    write_deliverables(project_id, 3, tmp_path)
    cell = openpyxl.load_workbook(tmp_path / "bom-v3.xlsx")["材料清单"]["B2"]
    assert cell.value == "=1+1" and cell.data_type == "s"


@pytest.mark.parametrize("tampered", [False, True])
def test_snapshot_only_copies_exact_manifest_binaries_with_matching_sha(tmp_path, tampered):
    project = tmp_path / "projects/demo"
    folder = project / "exports/v1"
    folder.mkdir(parents=True)
    exporter.write_json(project / "designs/v1.json", {"version": 1})
    names = ["layout-v1.dxf", "bom-v1.xlsx", "layout-v2.dxf", "secret.bin"]
    content = b"\xff\x00not JSON\n"
    files = {name: hashlib.sha256(content).hexdigest() for name in names}
    for name in names:
        (folder / name).write_bytes(content)
    files["../escape.xlsx"] = files[names[0]]
    if tampered:
        (folder / names[0]).write_bytes(b"tampered")
    exporter.write_json(folder / "manifest.json", {"version": 1, "files": files})
    events = [dict(seq=1, type="tool/call", data=dict(name="gz_export", toolUseId="export",
                                                     args=dict(project_id="demo", version=1)))]
    anomalies = []
    copied = exporter.snapshot(events, "demo", tmp_path / "run", tmp_path / "projects", anomalies)
    expected = ["bom-v1.xlsx"] if tampered else names[:2]
    binaries = {Path(name).name for name in copied if not name.endswith(".json")}
    assert binaries == set(expected)
    for name in expected:
        assert (tmp_path / "run/project/exports/v1" / name).read_bytes() == content
        assert copied["project/exports/v1/" + name] == files[name]
    assert bool(anomalies) == tampered


def test_snapshot_refuses_binary_resolving_outside_export_directory(tmp_path, monkeypatch):
    project = tmp_path / "projects/demo"
    folder = project / "exports/v1"
    folder.mkdir(parents=True)
    exporter.write_json(project / "designs/v1.json", {"version": 1})
    outside = tmp_path / "outside.xlsx"
    outside.write_bytes(b"not a deliverable")
    linked = folder / "bom-v1.xlsx"
    linked.write_bytes(outside.read_bytes())
    original_resolve = Path.resolve
    monkeypatch.setattr(Path, "resolve", lambda path, *args, **kwargs:
                        outside if path == linked else original_resolve(path, *args, **kwargs))
    exporter.write_json(folder / "manifest.json", {"version": 1, "files": {"bom-v1.xlsx": hashlib.sha256(outside.read_bytes()).hexdigest()}})
    events = [dict(seq=1, type="tool/call", data=dict(name="gz_export", toolUseId="export", args=dict(project_id="demo", version=1)))]
    anomalies = []
    copied = exporter.snapshot(events, "demo", tmp_path / "run", tmp_path / "projects", anomalies)
    assert anomalies and not any(path.endswith(".xlsx") for path in copied)
