"""Offline production tools, real rays and committed face examples."""
import json
import math
import os
import hashlib
from pathlib import Path
from shutil import copyfile
from unittest.mock import patch

import numpy as np
import pandas as pd
import pytest
import requests

from guangzhu import analysis, brief, layout, mcp_actions, project as storage, shading, roofplane
from guangzhu.config import ROOT
from guangzhu.models import Design, LayoutParams, Project
from guangzhu.shading_geometry import build_scene, ray_hits, ray_candidates, sun_vector
from guangzhu.solar import sun_position
from test_t32_engineering import project


def freeze_demo(assets, template, identifier, version):
    """Actual outputs for the next task's before capture; never an expected-value oracle."""
    from scripts.baseline import canonical
    target = assets / 'demo-baseline' / template
    manifest = {}
    for label, value in [('generated', 1), ('final', version)]:
        directory = target / label
        directory.mkdir(parents=True, exist_ok=True)
        paths = dict(project='project.json', layout=f'designs/v{value}.json')
        paths.update({name: f'analysis/v{value}/{name}.json' for name in ('summary', 'electrical', 'economics', 'structure', 'rules')})
        for name, relative in paths.items():
            raw = storage._path(identifier, relative).read_bytes()
            (directory / (name + '.json')).write_bytes(raw)
            stable = json.dumps(canonical(json.loads(raw)), ensure_ascii=False, indent=2) + '\n'
            (directory / (name + '-canonical.json')).write_text(stable, encoding='utf-8', newline='\n')
            manifest[f'{label}/{name}.json'] = hashlib.sha256(raw).hexdigest()
    project = storage.load(identifier)
    info = dict(template=template, project_id=identifier, generated_version=1, final_version=version,
                generated_params=storage.load_design(identifier, 1).params.model_dump(),
                final_params=storage.load_design(identifier, version).params.model_dump(),
                seam_setback_m=project.roof.seam_setback_m, raw_sha256=manifest,
                note='当前真实离线输出；canonical只省略时间/耗时供累积回归，raw保留原始字节；不是独立数学期望。')
    (target / 'manifest.json').write_text(json.dumps(info, ensure_ascii=False, indent=2), encoding='utf-8')


@pytest.mark.parametrize('rotation', [0, 37])
def test_saw_analytic_rays_and_real_winter_morning(rotation):
    from shapely.affinity import rotate
    from shapely.geometry import Polygon
    from guangzhu.models import Roof
    record = project('sawtooth', teeth=4, slope_deg=10)
    record.roof.outline = [(0, 0), (100, 0), (100, 40), (0, 40)]
    record.roof.planes = [type(record.roof.planes[0]).model_validate(p) for p in roofplane.generate(
        record.roof.outline, dict(kind='sawtooth', teeth=4, slope_deg=10))['planes']]
    design = Design(params=LayoutParams())
    scene = build_scene(record, design)
    points = np.array([[50, 12, .4526539614169302], [50, 18, 1.51061584566772]])
    if rotation:
        data = record.roof.model_dump()
        data['outline'] = list(rotate(Polygon(data['outline']), rotation, origin=(0, 0)).exterior.coords)[:-1]
        for plane in data['planes']:
            plane['outline'] = list(rotate(Polygon(plane['outline']), rotation, origin=(0, 0)).exterior.coords)[:-1]
            plane['slope_azimuth_deg'] = (plane['slope_azimuth_deg'] - rotation) % 360
        record.roof = Roof.model_validate(data)
        angle = math.radians(rotation)
        points[:, :2] = points[:, :2] @ np.array([[math.cos(angle), math.sin(angle)], [-math.sin(angle), math.cos(angle)]])
        scene = build_scene(record, design)
    vector = sun_vector(10, 125 - rotation)
    normal = sun_vector(80, 180 - rotation)
    assert np.dot(normal, vector) > 0  # Actual front incidence, independent from monthly energy-bin weights.
    assert list(ray_hits(scene, points, vector, np.array([-1, -1]))) == ['roof', '']
    point_index, face_index = ray_candidates(scene, points, vector, np.array([-1, -1]))
    assert scene.identifiers[face_index[0]] == 'P2'
    if not rotation:
        actual = sun_position(pd.DatetimeIndex(['2025-12-21T00:00:00Z']), record.site.lat, record.site.lon).iloc[0]
        assert actual.elevation > 0
        assert ray_hits(scene, points[:1], sun_vector(actual.elevation, actual.azimuth), np.array([-1]))[0] == 'roof'


@pytest.mark.parametrize('kind', ['gable', 'hip'])
def test_convex_roof_has_no_manufactured_front_beam_loss_and_net_faces_sum(kind):
    record = project(kind)
    design = layout.generate_layout(record, {})
    shadows = shading.analyze_design(record, design, offline=True)
    assert shadows['summary']['source_loss_pct']['roof'] == pytest.approx(0, abs=1e-9)
    energy, _ = analysis.generation(record, design, shadows, design.statistics['dc_kw'], offline=True)
    assert sum(p['p50'] for p in energy['planes']) == pytest.approx(energy['p50'])
    assert sum(p['dc_kw'] for p in energy['planes']) == pytest.approx(design.statistics['dc_kw'])


def test_net_yield_selection_is_inclusive_and_ignores_empty_faces():
    from guangzhu.pitched_compare import eligible
    assert eligible([dict(id=1, dc_kw=1, per_kw=1000), dict(id=2, dc_kw=2, per_kw=950),
                     dict(id=3, dc_kw=2, per_kw=949.999), dict(id=4, dc_kw=0, per_kw=2000)]) == [1, 2]


def test_final_fast_candidate_rechecks_real_net_yield_after_rule_repairs(tmp_path, monkeypatch):
    from guangzhu.models import RoofPlane, Obstacle, StructuredRequirements
    from guangzhu.pitched_compare import eligible
    monkeypatch.setattr(storage, 'PROJECTS_DIR', tmp_path / 'projects')
    monkeypatch.setattr(requests.Session, 'request', lambda *a, **k: pytest.fail('T32 offline comparison'))
    storage.init(ROOT / 'data/demo/projects/changzhou-demo', 't32-converge')
    record = project('hip', slope_deg=6)
    record.id = 't32-converge'
    record.roof.outline = [(0, 0), (60, 0), (60, 40), (0, 40)]
    record.roof.planes = [RoofPlane.model_validate(p) for p in roofplane.generate(record.roof.outline,
                         dict(kind='hip', slope_deg=6))['planes']]
    record.roof.load_capacity_kn_m2 = 3
    record.roof.obstacles = [Obstacle(id='device', label='局部示意设备', kind='prism',
                                     footprint=[[5, 18], [7, 18], [7, 20], [5, 20]], height_m=3)]
    # 420000 before costs included the 0.70 yuan/W calibrated to the industry average; scaled by 2.59/1.89 the budget binds the same way
    record.requirements.structured = StructuredRequirements(budget_cny_max=575000)
    storage.save(record)
    mcp_actions._generate_layout(record.id, {}, offline=True)
    mcp_actions._analyze(record.id, 1, offline=True)
    result = mcp_actions._compare_options(record.id, 1, offline=True)
    selection = result['plane_selection']
    planes = lambda version: mcp_actions._read(record.id, f'analysis/v{version}/economics.json')['analysis_metadata']['yield_kwh']['planes']
    initial = planes(selection['initial_basis_version'])
    assert eligible(initial) == [2, 3, 4]
    assert selection['attempts'] == 2 and selection['selected'] == [2, 3]
    assert result['options'][1]['version'] == selection['basis_version']
    assert result['options'][1]['planes'] == eligible(planes(selection['basis_version'])) == [2, 3]
    # Real S01 repair improves P4's net yield but leaves it below 95%; it must be removed and reanalysed.
    repaired = planes(selection['basis_version'] - 1)
    before_repair = planes(selection['basis_version'] - 2)
    yield4 = lambda rows: next(p['per_kw'] for p in rows if p['id'] == 4)
    assert yield4(initial) >= .95 * max(p['per_kw'] for p in initial)
    assert yield4(before_repair) < yield4(repaired) < .95 * max(p['per_kw'] for p in repaired)
    metric = result['options'][1]
    final = mcp_actions._analyze(record.id, metric['version'], offline=True)
    assert metric['hard_pass'] and final['rules']['hard_failed'] == 0
    assert metric['annual_kwh'] == final['yield_kwh']['p50']
    assert metric['capex_cny'] == final['economics']['capex_cny']
    assert result['recommended'] == mcp_actions.recommend_options(result['current'], result['options'][:2])
    destination = ROOT / 'reports/assets/T32' if os.environ.get('GUANGZHU_T32_CAPTURE') == '1' else tmp_path
    destination.mkdir(parents=True, exist_ok=True)
    evidence = dict(compare=result, initial=initial, before_repair=before_repair, repaired=repaired,
                    final_planes=planes(selection['basis_version']), final_analysis=final)
    (destination / 'net-yield-convergence.json').write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding='utf-8')


@pytest.fixture(scope='module')
def pitched_flow(tmp_path_factory):
    temporary = tmp_path_factory.mktemp('t32-flow')
    assets = ROOT / 'reports/assets/T32' if os.environ.get('GUANGZHU_T32_CAPTURE') == '1' else temporary / 'assets'
    assets.mkdir(parents=True, exist_ok=True)
    with patch.object(storage, 'PROJECTS_DIR', temporary / 'projects'):
        with patch.object(requests.Session, 'request', side_effect=AssertionError('T32离线流程禁止联网')):
            created = mcp_actions._new_project('t32-offline', 'pitched-gable-demo')
            assert created['roof_type'] == 'pitched' and created['plane_count'] == 2
            mcp_actions._set_requirements('t32-offline', '整栋双坡离线示意', {})
            generated = mcp_actions._generate_layout('t32-offline', {}, offline=True)
            assert generated['mount'] == 'plane_flush'
            summary = mcp_actions._analyze('t32-offline', 1, offline=True)
            compare = mcp_actions._compare_options('t32-offline', 1, offline=True)
            version = compare['recommended']['balanced']['version']
            summary = mcp_actions._analyze('t32-offline', version, offline=True)
            context = dict(project=storage.load('t32-offline').model_dump(by_alias=True),
                           design=storage.load_design('t32-offline', version).model_dump(mode='json'),
                           version=version, summary=summary, compare=compare)
            for key in ('rules', 'economics', 'electrical', 'shading', 'structure'):
                context[key] = mcp_actions._read('t32-offline', f'analysis/v{version}/{key}.json')
            assert context['rules']['hard_failed'] == 0
            assert all(r['status'] == 'na' for r in context['rules']['items'] if r['id'] in ('G02', 'G05'))
            assert len([r for r in context['rules']['items'] if r['id'] == 'G01']) == 2
            assert len([r for r in context['rules']['items'] if r['id'] == 'L01']) == 2
            for string in context['electrical']['strings']:
                assert len({mid.split('-')[0] for mid in string['module_ids']}) == 1
            crosscheck = mcp_actions._crosscheck_yield('t32-offline', version, offline=True)
            exported = mcp_actions._export('t32-offline', version, offline=True)
            assert exported['ok']
            context.update(crosscheck=crosscheck, review=None, approval=None, sessions=[], verification={}, sources={},
                           versions=[], relaxations=[], history=[], run={}, requirements=[dict(text='整栋双坡离线示意')],
                           generated_at='2026-10-06T00:00:00+08:00', run_id='T32-offline-no-AI', project_id='t32-offline')
            from guangzhu.report.view import render
            contents = render(context)
            assert 'P1 坡面' in contents and 'P2 坡面' in contents and '每千瓦净年发电' in contents
            assert '未提供 AI 会话' in contents
            (assets / 'proposal.html').write_text(contents, encoding='utf-8')
            scene = dict(project=context['project'], design=context['design'], shading=context['shading'],
                         electrical=context['electrical'], wind=context['structure']['wind'],
                         roof_geometry=roofplane.scene_geometry(storage.load('t32-offline').roof))
            for filename, value in [('context.json', context), ('scene.json', scene),
                                    ('tool-flow.json', dict(created=created, layout=generated, analysis=summary,
                                                           compare=compare, crosscheck=crosscheck, export=exported))]:
                (assets / filename).write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')
            for filename in [*exported['files'], 'manifest.json']:
                copyfile(storage._path('t32-offline', f'exports/v{version}/{filename}'), assets / filename)
            freeze_demo(assets, 'pitched-gable-demo', 't32-offline', version)
            yield dict(assets=assets, context=context, scene=scene, temporary=temporary, version=version)


def test_saw_demo_normal_tools_final_analysis_and_export(tmp_path, monkeypatch):
    monkeypatch.setattr(storage, 'PROJECTS_DIR', tmp_path / 'projects')
    monkeypatch.setattr(requests.Session, 'request', lambda *a, **k: pytest.fail('T32 saw offline tools'))
    target = ROOT / 'reports/assets/T32' if os.environ.get('GUANGZHU_T32_CAPTURE') == '1' else tmp_path / 'assets'
    target.mkdir(parents=True, exist_ok=True)
    created = mcp_actions._new_project('t32-saw', 'pitched-saw-demo')
    mcp_actions._set_requirements('t32-saw', '四齿整栋离线示意', {})
    generated = mcp_actions._generate_layout('t32-saw', {}, offline=True)
    initial = mcp_actions._analyze('t32-saw', 1, offline=True)
    compare = mcp_actions._compare_options('t32-saw', 1, offline=True)
    version = compare['recommended']['balanced']['version']
    summary = mcp_actions._analyze('t32-saw', version, offline=True)
    assert summary['rules']['hard_failed'] == 0
    exported = mcp_actions._export('t32-saw', version, offline=True)
    assert exported['ok'] and created['plane_count'] == 8
    design = storage.load_design('t32-saw', version)
    assert all(int(m.id.split('-')[0][1:]) in (1, 3, 5, 7) for m in design.modules)
    freeze_demo(target, 'pitched-saw-demo', 't32-saw', version)
    data = dict(created=created, generated=generated, initial=initial, compare=compare, final=summary, exported=exported)
    (target / 'saw-tool-flow.json').write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    for filename in [*exported['files'], 'manifest.json']:
        destination = target / 'saw-export' / filename
        destination.parent.mkdir(parents=True, exist_ok=True)
        copyfile(storage._path('t32-saw', f'exports/v{version}/{filename}'), destination)


def test_production_flow_pitched_brief_and_actual_bom(pitched_flow):
    from openpyxl import load_workbook
    import ezdxf
    data = pitched_flow
    result = data['context']
    assert '"compare":true' in brief.render_brief(dict(template='pitched-gable-demo', id='t32-brief', text='整栋示意', compare=True))
    book = load_workbook(data['assets'] / f"bom-v{data['version']}.xlsx", read_only=True)
    assert {'坡面', '坡面造价', '造价明细'} <= set(book.sheetnames)
    faces = list(book['坡面'].values)[1:]
    assert sum(r[8] for r in faces) == len(result['design']['modules'])
    assert sum(r[10] for r in faces) == pytest.approx(result['economics']['analysis_metadata']['yield_kwh']['p50'])
    prices = list(book['坡面造价'].values)[1:]
    assert math.fsum(r[5] for r in prices) == pytest.approx(result['economics']['capex_cny'])
    book.close()
    drawing = ezdxf.readfile(data['assets'] / f"layout-v{data['version']}.dxf")
    assert {'P1_metal', 'P2_metal'} <= {layer.dxf.name for layer in drawing.layers}
    assert len(analysis.serialized(result['summary']).encode()) <= 1920
    assert len(analysis.serialized(result['compare']).encode()) <= 1920
