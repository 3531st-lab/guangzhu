"""Offline T30 numeric evidence and real exports in an explicitly owned directory."""
import argparse
import json
from datetime import datetime, timedelta, timezone
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'backend'))
from guangzhu import analysis, economics, finance, layout, project as storage, tariff
from guangzhu.models import Economics, SiteFacts
from guangzhu.report import render_pdf
from guangzhu.report.deliverables import write_deliverables
from guangzhu.report.view import render


def write(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + '\n', encoding='utf-8')


def energy(date, amount=24):
    stamp = datetime.fromisoformat(date).replace(tzinfo=timezone(timedelta(hours=8)))
    return dict(p50=amount, low=amount, high=amount, hourly=dict(kwh=[amount / 24] * 24,
        utc=[(stamp + timedelta(hours=h)).astimezone(timezone.utc).isoformat() for h in range(24)]))


def numeric(project):
    days = []
    for date, kva in [('2023-03-15', 100), ('2023-07-15', 315), ('2023-01-23', 315)]:
        project.economics = Economics(self_use_ratio=1)
        project.facts = SiteFacts(transformer_kva=kva)
        days.append(dict(date=date, result=tariff.integrate(project, energy(date), 1, .65)))
    ledgers = {mode: finance.ledger(11300, 11300, 1000, 2260, 100, mode=mode, rent=300,
                                  replacement=1130, discount=.08) for mode in ('self', 'emc', 'rent')}
    ledgers['self_loan'] = finance.ledger(11300, 11300, 1000, 2260, 100, replacement=1130,
                                        loan_ratio=.5, loan_rate=.05, loan_years=20, discount=.08)
    exports = []
    for allowed in (True, False):
        project.facts = None
        export_price = 0 if allowed else .391
        project.economics = Economics(price_export_cny_kwh=export_price, self_use_ratio=.8)
        project.requirements.structured.export_allowed = allowed
        result = economics.calculate(project, 100, energy('2023-03-15', 100000), sensitivity=False)
        exports.append(dict(export_allowed=allowed, export_price_cny_kwh=export_price,
                            self_use=result['self_use'], certificates=result['certificates'],
                            income=result['income_breakdown'], accounts=result['accounts']))
    return dict(entrypoints='tariff.integrate; finance.ledger; economics.calculate', days=days, ledgers=ledgers, exports=exports,
                reference_tax=finance.ledger(11300, 11300, 1000, 2260, 100),
                reference_debt=finance.debt_schedule(100000, .5, .05, 20, 25),
                zero_interest=finance.debt_schedule(10000, .5, 0, 5, 25))


def comparison(output):
    rows = []
    for path in sorted((output / 'before').glob('*-economics.json')):
        name = path.name.removesuffix('-economics.json')
        old = json.loads(path.read_text(encoding='utf-8'))
        new = json.loads((output / 'after' / path.name).read_text(encoding='utf-8'))
        params = new['params']
        p50 = new['cashflows'][1]['generation_kwh']
        ratio = params['self_use_ratio']
        flat_income = p50 * (ratio * params['price_self_cny_kwh'] + (1 - ratio) * params['price_export_cny_kwh'])
        base_om = new['operating_costs']['base_om_cny']
        cost_only = finance.ledger(new['capex_cny'], 0, flat_income, 0, base_om, degradation=params['degradation_annual_pct'])
        tou_income = new['income_breakdown']['self_saving_cny'] + new['income_breakdown']['export_cny']
        cost_tou = finance.ledger(new['capex_cny'], 0, tou_income, 0, base_om, degradation=params['degradation_annual_pct'])
        rows.append(dict(roof=name, before_capex=old['capex_cny'], after_capex=new['capex_cny'],
                         before_price=old['params']['price_self_cny_kwh'], after_price=new['tariff']['weighted_price_cny_kwh'],
                         before_payback=old['payback_years'], cost_only_payback=cost_only['pretax']['payback_years'],
                         cost_and_tou_payback=cost_tou['pretax']['payback_years'], after_payback=new['payback_years'],
                         after_posttax=new['posttax']['payback_years']))
    write(output / 'six-demo-comparison.json', rows)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--work', type=Path, required=True)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    storage.PROJECTS_DIR = args.work / 'projects'
    identifier = 't30-offline-export'
    instance = storage.load(identifier) if storage._path(identifier).exists() else storage.init(
        ROOT / 'data/demo/projects/changzhou-demo', identifier)
    write(args.output / 'numeric-actual.json', numeric(instance.model_copy(deep=True)))
    instance.economics.loan_ratio = .5
    instance.facts = SiteFacts(transformer_kva=315, roof_age_years=10)
    storage.save(instance)
    if storage._path(instance.id, 'designs/v1.json').exists():
        design = storage.load_design(instance.id, 1)
    else:
        design = layout.generate_layout(instance, {})
        storage.save_design(instance.id, design)
    summary = analysis.analyze(instance.id, design.version, offline=True)
    context = dict(project=instance.model_dump(mode='json'), design=design.model_dump(mode='json'), summary=summary,
                   version=design.version, compare=None, facts=instance.facts.model_dump(), review=None, approval=None,
                   sessions=[], requirements=[], verification={}, sources={}, versions=[], relaxations=[], history=[], run={},
                   generated_at='2026-10-06T00:00:00+08:00', run_id='T30-offline-no-AI', project_id=instance.id, crosscheck=None)
    for name in ('electrical', 'shading', 'structure', 'rules', 'economics'):
        path = storage._path(instance.id, f'analysis/v{design.version}/{name}.json')
        context[name] = json.loads(path.read_text(encoding='utf-8'))
    folder = args.output / 'export'
    folder.mkdir(exist_ok=True)
    write(folder / 'context.json', context)
    html = folder / 'proposal.html'
    html.write_text(render(context), encoding='utf-8')
    state = render_pdf(html.resolve(), dict(files=[]))
    write(folder / 'pdf-state.json', state)
    if state['status'] != 'ready':
        raise RuntimeError('Real offline PDF was not created')
    write_deliverables(instance.id, design.version, folder)
    import pypdfium2 as pdfium
    from openpyxl import load_workbook
    book = load_workbook(folder / f'bom-v{design.version}.xlsx', data_only=True)
    pdf = pdfium.PdfDocument(folder / 'proposal.pdf')
    texts = [page.get_textpage().get_text_range() for page in pdf]
    for title in ('造价明细', '电价与收益', '税和贷款'):
        page = next(i for i, text in enumerate(texts) if title in text)
        pdf[page].render(scale=1.3).to_pil().save(folder / (title + '.png'))
    write(folder / 'reopened.json', dict(pdf_pages=len(pdf), sheets=book.sheetnames,
                                       cost_total=book['造价明细'].cell(book['造价明细'].max_row, 5).value))
    comparison(args.output)
    print(json.dumps(dict(pdf_pages=len(pdf), sheets=book.sheetnames), ensure_ascii=False), flush=True)


if __name__ == '__main__':
    main()
